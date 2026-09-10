#!/usr/bin/env python3
# ruff: noqa: ANN401, C901, CPY001, D103, DOC201, E501, I001, PIE808, PLR0912, PLR0914, PLR0915, PLR0916, PLR2004, SLF001, T201
"""Build the advisor-facing final experiment report from accepted local artifacts."""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import math
import re
from itertools import pairwise
from pathlib import Path
from statistics import median
from typing import Any

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont


BLUE = "2E6F9E"
ORANGE = "D97934"
INK = "1F2933"
MUTED = "52606D"
LIGHT = "E8EDF2"
NORMAL_LABEL = "VAE no equivariante"
SO2_LABEL = "VAE equivariante SO(2)"
HEATMAP_LIGHT_TEXT_THRESHOLD = 0.55
FIXED25_COUNT = 25
FIRST_ORBIT_SHEET_COUNT = 15
ROTATION_GEOMETRY_SUMMARY_SHA256 = (
    "aab68ca5b7cbe989ea1cd79740a64d83784441f85ffde09a6ee6ebebe7faf43c"
)
ROTATION_GEOMETRY_ADDENDUM_SHA256 = (
    "ec74b37f509205bca12294321844a3cfb48dbd2f1763c8ff94290fe064e86f94"
)
ROTATION_GEOMETRY_SPEC_SHA256 = (
    "fbc58459d214ee32f1148f6b58507820586be1993749811294f14eddbf467b61"
)
ROTATION_GEOMETRY_CONTRACT_SHA256 = (
    "1dc6975979b120d85680526a3177e30caf93ca2eba7889a4452f2e78ca399756"
)
DECODED_TRANSFORM_SUMMARY_SHA256 = (
    "616e3a6c9caa6b8e72a0b21f41a5121d7a5baa420f8023b35775fedaf252a754"
)
DECODED_TRANSFORM_MANIFEST_SHA256 = (
    "a1168b419080305138f7aedca4c1ed345c3d5947e78299b0e5ea1430a75d3c50"
)
DECODED_TRANSFORM_ADDENDUM_SHA256 = (
    "b0fa0cc2eebe1207fcc2da7a00c30b9b85c40f4353e332131423f76f6c380e34"
)
DECODED_TRANSFORM_SPEC_SHA256 = (
    "7a472a2de56813544506b4a9eb58e2f1f152fcee8f7d60e65ed2c2793b349679"
)
DECODED_TRANSFORM_CONTRACT_SHA256 = (
    "905ef933a51c26bb3f01fcef4bb4985fa98e00f1dbecd814da3df488bd7c1018"
)


SPANISH_REPLACEMENTS = {
    "Aun": "Aún",
    "Area": "Área",
    "Atribucion": "Atribución",
    "Comparacion": "Comparación",
    "Conclusion": "Conclusión",
    "Cuadricula": "Cuadrícula",
    "Descripcion": "Descripción",
    "Diagnostico": "Diagnóstico",
    "Donde": "Dónde",
    "Entropia": "Entropía",
    "Definicion": "Definición",
    "Desviacion": "Desviación",
    "Dimension": "Dimensión",
    "Direccion": "Dirección",
    "Grafica": "Gráfica",
    "Graficas": "Gráficas",
    "Interpretacion": "Interpretación",
    "Implementacion": "Implementación",
    "Metricas": "Métricas",
    "Metrica": "Métrica",
    "Parametros": "Parámetros",
    "Poblacion": "Población",
    "Orbita": "Órbita",
    "Perdida": "Pérdida",
    "Reconstruccion": "Reconstrucción",
    "Razon": "Razón",
    "Relacion": "Relación",
    "Seccion": "Sección",
    "Senal": "Señal",
    "Senales": "Señales",
    "Si": "Sí",
    "Simetria": "Simetría",
    "Sintesis": "Síntesis",
    "Tamano": "Tamaño",
    "Validacion": "Validación",
    "Visualizacion": "Visualización",
    "amplio": "amplió",
    "agregacion": "agregación",
    "analisis": "análisis",
    "almaceno": "almacenó",
    "angulos": "ángulos",
    "aritmetica": "aritmética",
    "acompana": "acompaña",
    "aplico": "aplicó",
    "automaticamente": "automáticamente",
    "biologica": "biológica",
    "area": "área",
    "arquitectonica": "arquitectónica",
    "arquitectonico": "arquitectónico",
    "asi": "así",
    "atencion": "atención",
    "atribucion": "atribución",
    "clasificacion": "clasificación",
    "ciclico": "cíclico",
    "ciclicos": "cíclicos",
    "compactacion": "compactación",
    "comparacion": "comparación",
    "compensacion": "compensación",
    "conclusion": "conclusión",
    "contribucion": "contribución",
    "cientificamente": "científicamente",
    "cuadratico": "cuadrático",
    "diagnostico": "diagnóstico",
    "diagnosticos": "diagnósticos",
    "descripcion": "descripción",
    "deberia": "debería",
    "despues": "después",
    "desviacion": "desviación",
    "direccion": "dirección",
    "dispersion": "dispersión",
    "entropia": "entropía",
    "estandar": "estándar",
    "establecio": "estableció",
    "estan": "están",
    "estratifico": "estratificó",
    "estadistica": "estadística",
    "evaluacion": "evaluación",
    "evaluo": "evaluó",
    "explicacion": "explicación",
    "explicita": "explícita",
    "especifica": "específica",
    "especificas": "específicas",
    "entreno": "entrenó",
    "experimentacion": "experimentación",
    "fragiles": "frágiles",
    "fluctuan": "fluctúan",
    "grafica": "gráfica",
    "graficas": "gráficas",
    "histopatologia": "histopatología",
    "histopatologico": "histopatológico",
    "guardo": "guardó",
    "incluyo": "incluyó",
    "inspeccion": "inspección",
    "implementacion": "implementación",
    "interpretacion": "interpretación",
    "laminas": "láminas",
    "limitacion": "limitación",
    "linea": "línea",
    "maximo": "máximo",
    "mas": "más",
    "mejoro": "mejoró",
    "metrica": "métrica",
    "metricas": "métricas",
    "mostro": "mostró",
    "ningun": "ningún",
    "numero": "número",
    "numerica": "numérica",
    "normalizacion": "normalización",
    "numericamente": "numéricamente",
    "observo": "observó",
    "organizacion": "organización",
    "oraculo": "oráculo",
    "orbita": "órbita",
    "parametrica": "paramétrica",
    "parametros": "parámetros",
    "patron": "patrón",
    "pequena": "pequeña",
    "pequeno": "pequeño",
    "perdida": "pérdida",
    "periodica": "periódica",
    "pixel": "píxel",
    "pixeles": "píxeles",
    "poblacion": "población",
    "proposito": "propósito",
    "probabilistica": "probabilística",
    "presentacion": "presentación",
    "prediccion": "predicción",
    "reconstruccion": "reconstrucción",
    "representacion": "representación",
    "recupero": "recuperó",
    "razon": "razón",
    "resumenes": "resúmenes",
    "reduccion": "reducción",
    "regimenes": "regímenes",
    "restriccion": "restricción",
    "rotacion": "rotación",
    "segun": "según",
    "seleccion": "selección",
    "senal": "señal",
    "senales": "señales",
    "simetria": "simetría",
    "simultaneo": "simultáneo",
    "simultaneos": "simultáneos",
    "simultaneamente": "simultáneamente",
    "solicito": "solicitó",
    "sintesis": "síntesis",
    "estadisticas": "estadísticas",
    "tamano": "tamaño",
    "tambien": "también",
    "unica": "única",
    "unico": "único",
    "uso": "usó",
    "validacion": "validación",
    "utiles": "útiles",
    "variacion": "variación",
    "visualizacion": "visualización",
    "verificacion": "verificación",
}


def es(text: str) -> str:
    """Apply conservative Spanish accent restoration to display text."""
    for source, target in SPANISH_REPLACEMENTS.items():
        text = re.sub(rf"\b{re.escape(source)}\b", target, text)
    return text


def read_json(path: Path) -> dict[str, Any]:
    return json.loads(path.read_text(encoding="utf-8"))


def read_csv(path: Path) -> list[dict[str, str]]:
    with path.open(encoding="utf-8", newline="") as handle:
        return list(csv.DictReader(handle))


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def validate_rotation_geometry(
    document: dict[str, Any],
    addendum: dict[str, Any],
) -> dict[str, Any]:
    """Reject stale, incomplete, or overinterpreted Spec 0050 evidence.

    Raises:
        TypeError: A required mapping has the wrong type.
        ValueError: An identity, decision, median, or sampling unit is inconsistent.

    """
    if document.get("schema") != "spec0050.corrected_rotation_geometry.v1":
        message = "La evidencia geométrica corregida usa un esquema inesperado"
        raise ValueError(message)
    if document.get("spec_sha256") != ROTATION_GEOMETRY_SPEC_SHA256:
        message = "La evidencia geométrica no coincide con el spec de lanzamiento"
        raise ValueError(message)
    if document.get("contract_sha256") != ROTATION_GEOMETRY_CONTRACT_SHA256:
        message = "La evidencia geométrica no coincide con el contrato bloqueado"
        raise ValueError(message)
    contract = document.get("contract")
    if not isinstance(contract, dict):
        message = "Falta el contrato embebido de Spec 0050"
        raise TypeError(message)
    scope = contract.get("scope")
    if (
        not isinstance(scope, dict)
        or scope.get("population") != "fixed_validation_25"
        or scope.get("validation_only") is not True
        or scope.get("sealed_test_access") is not False
        or scope.get("training_allowed") is not False
    ):
        message = "El alcance del artefacto geométrico no es validación fija aislada"
        raise ValueError(message)
    tensor_contract = contract.get("tensor_contract")
    if not isinstance(tensor_contract, dict) or tensor_contract.get(
        "angles_degrees",
    ) != {"start": 0, "step": 1, "stop_inclusive": 359}:
        message = "El barrido corregido no contiene exactamente 0°--359°"
        raise ValueError(message)

    comparisons = document.get("comparisons")
    branches = document.get("branches")
    decisions = document.get("decisions")
    if not all(
        isinstance(section, dict) for section in (comparisons, branches, decisions)
    ):
        message = "Faltan secciones geométricas obligatorias"
        raise TypeError(message)
    for metric_name in (
        "local_linearity_ratio",
        "step_size_cv",
        "path_length",
        "curvature_median",
        "curvature_q90",
    ):
        metric = comparisons.get(metric_name)
        if not isinstance(metric, dict):
            message = f"Falta la métrica corregida {metric_name}"
            raise TypeError(message)
        for branch in ("normal", "so2"):
            values = metric.get(branch)
            stored = metric.get(f"{branch}_median")
            if not isinstance(values, list) or len(values) != FIXED25_COUNT:
                message = f"{metric_name}.{branch} no contiene 25 valores"
                raise ValueError(message)
            numeric = [float(value) for value in values]
            if any(not math.isfinite(value) for value in numeric):
                message = f"{metric_name}.{branch} contiene valores no finitos"
                raise ValueError(message)
            if not isinstance(stored, (int, float)) or not math.isclose(
                float(stored),
                float(median(numeric)),
                rel_tol=0.0,
                abs_tol=1e-12,
            ):
                message = f"Mediana inconsistente para {metric_name}.{branch}"
                raise ValueError(message)

    split = contract.get("split")
    patch_hashes = contract.get("inputs", {}).get("patch_sha256")
    if (
        not isinstance(split, dict)
        or len(split.get("fit_ranks", [])) != 17
        or len(split.get("heldout_ranks", [])) != 8
        or not isinstance(patch_hashes, list)
        or len(patch_hashes) != FIXED25_COUNT
        or len(set(patch_hashes)) != FIXED25_COUNT
    ):
        message = "La identidad o la partición determinista fixed-25 es inválida"
        raise ValueError(message)
    if document.get("selected_dimension", {}).get("dimension") != 1:
        message = "El artefacto aceptado ya no contiene la dimensión nula revisada"
        raise ValueError(message)
    if (
        comparisons.get("h1_passes") is not False
        or comparisons.get("h2_passes") is not False
        or decisions.get("f1_h3_passes") is not False
        or decisions.get("shared_action", {}).get("so2_specific_advantage") is not False
        or decisions.get("learned_factorization", {}).get("so2_vae", {}).get("status")
        != "unresolved"
    ):
        message = "Las decisiones geométricas no coinciden con la revisión aceptada"
        raise ValueError(message)

    if addendum.get("schema") != "spec0050.rotation_geometry_review_addendum.v1":
        message = "El addendum de revisión usa un esquema inesperado"
        raise ValueError(message)
    if addendum.get("selected_dimension") != 1 or addendum.get(
        "heldout_ranks",
    ) != split.get("heldout_ranks"):
        message = "El addendum no corresponde al split o dimensión aceptados"
        raise ValueError(message)
    for branch in ("normal_vae", "so2_vae"):
        rows = addendum.get(branch, {}).get("per_patch")
        if not isinstance(rows, list) or len(rows) != 8:
            message = f"El addendum {branch} no conserva ocho unidades de muestreo"
            raise ValueError(message)
        if [row.get("rank") for row in rows] != split.get("heldout_ranks"):
            message = f"Los rangos retenidos del addendum {branch} no coinciden"
            raise ValueError(message)

    return document


def validate_decoded_transform(
    document: dict[str, Any],
    addendum: dict[str, Any],
) -> dict[str, Any]:
    """Reject stale or overinterpreted Spec 0051 evidence.

    Raises:
        TypeError: A required section has the wrong type.
        ValueError: An identity, count, metric, or decision is inconsistent.

    """
    if document.get("schema") != "spec0051.decoded_latent_transform.v1":
        message = "La evidencia de decodificación usa un esquema inesperado"
        raise ValueError(message)
    if document.get("spec_sha256") != DECODED_TRANSFORM_SPEC_SHA256:
        message = "La evidencia de decodificación no coincide con Spec 0051"
        raise ValueError(message)
    if document.get("contract_sha256") != DECODED_TRANSFORM_CONTRACT_SHA256:
        message = "La evidencia de decodificación no coincide con su contrato"
        raise ValueError(message)
    contract = document.get("contract")
    if not isinstance(contract, dict):
        message = "Falta el contrato embebido de Spec 0051"
        raise TypeError(message)
    scope = contract.get("scope")
    if (
        not isinstance(scope, dict)
        or scope.get("population") != "fixed_validation_25"
        or scope.get("validation_only") is not True
        or scope.get("sealed_test_access") is not False
        or scope.get("training_allowed") is not False
    ):
        message = "El alcance de Spec 0051 no es validación fija aislada"
        raise ValueError(message)
    dense = contract.get("dense_rotations")
    if not isinstance(dense, dict) or dense.get("angles_degrees") != list(
        range(0, 360, 5),
    ):
        message = "El barrido decodificado no contiene exactamente 0,5,...,355"
        raise ValueError(message)
    margins = contract.get("decision_margins")
    expected_margins = {
        "ratio_strong": 0.5,
        "ratio_patch_success": 0.75,
        "minimum_relative_so2_advantage": 0.2,
    }
    if not isinstance(margins, dict) or any(
        not isinstance(margins.get(name), (int, float))
        or not math.isclose(
            float(margins[name]),
            expected,
            rel_tol=0.0,
            abs_tol=0.0,
        )
        for name, expected in expected_margins.items()
    ):
        message = "Los umbrales absolutos o comparativos de Spec 0051 cambiaron"
        raise ValueError(message)

    comparisons = document.get("comparisons")
    branches = document.get("branches")
    decisions = document.get("decisions")
    if not all(
        isinstance(section, dict) for section in (comparisons, branches, decisions)
    ):
        message = "Faltan secciones obligatorias de Spec 0051"
        raise TypeError(message)
    for metric_name in ("action_ratio", "canonical_ratio", "input_commutation_ratio"):
        metric = comparisons.get(metric_name)
        if not isinstance(metric, dict):
            message = f"Falta la métrica decodificada {metric_name}"
            raise TypeError(message)
        for branch in ("normal", "so2"):
            values = metric.get(branch)
            stored = metric.get(f"{branch}_median")
            if not isinstance(values, list) or len(values) != FIXED25_COUNT:
                message = f"{metric_name}.{branch} no contiene 25 valores"
                raise ValueError(message)
            numeric = [float(value) for value in values]
            if any(not math.isfinite(value) for value in numeric):
                message = f"{metric_name}.{branch} contiene valores no finitos"
                raise ValueError(message)
            if not isinstance(stored, (int, float)) or not math.isclose(
                float(stored),
                float(median(numeric)),
                rel_tol=0.0,
                abs_tol=1e-12,
            ):
                message = f"Mediana inconsistente para {metric_name}.{branch}"
                raise ValueError(message)
        paired = metric.get("paired")
        expected_favorable = 0 if metric_name == "input_commutation_ratio" else 25
        if (
            not isinstance(paired, dict)
            or paired.get("n_patches") != FIXED25_COUNT
            or paired.get("n_clusters") != 16
            or paired.get("so2_favorable_patch_count") != expected_favorable
        ):
            message = f"Comparación pareada inconsistente para {metric_name}"
            raise ValueError(message)

    normal_population = branches.get("normal_vae", {}).get("population", {})
    so2_population = branches.get("so2_vae", {}).get("population", {})
    expected_quarter = {
        "normal_action": 0.7094527561437537,
        "so2_action": 0.0000021849270987967017,
        "normal_canonical": 0.7106375984701042,
        "so2_canonical": 0.08885407753921012,
    }
    observed_quarter = {
        "normal_action": normal_population.get("exact_quarter_action_ratio_median"),
        "so2_action": so2_population.get("exact_quarter_action_ratio_median"),
        "normal_canonical": normal_population.get(
            "exact_quarter_canonical_ratio_median",
        ),
        "so2_canonical": so2_population.get("exact_quarter_canonical_ratio_median"),
    }
    for name, expected in expected_quarter.items():
        observed = observed_quarter[name]
        if not isinstance(observed, (int, float)) or not math.isclose(
            float(observed),
            expected,
            rel_tol=0.0,
            abs_tol=1e-12,
        ):
            message = f"Control exacto C4 inconsistente: {name}"
            raise ValueError(message)

    if (
        decisions.get("H1_decoded_rotational_action", {}).get("supported") is not False
        or decisions.get("H2_decoded_rotational_canonicalization", {}).get(
            "supported",
        )
        is not False
        or decisions.get("H3_decoded_residual_suppression", {}).get("supported")
        is not False
        or decisions
        .get("H4_reflection_robustness", {})
        .get("flip_diag", {})
        .get("interpretation")
        != "so2_checkpoint_only_not_architecture_proof"
    ):
        message = "Las decisiones de Spec 0051 no coinciden con la revisión"
        raise ValueError(message)

    if addendum.get("schema") != "spec0051.decoded_transform_review_addendum.v1":
        message = "El addendum de Spec 0051 usa un esquema inesperado"
        raise ValueError(message)
    source = addendum.get("source")
    acceptance = addendum.get("acceptance")
    publication = addendum.get("publication_state")
    if not all(
        isinstance(section, dict) for section in (source, acceptance, publication)
    ):
        message = "El addendum de Spec 0051 está incompleto"
        raise TypeError(message)
    if (
        source.get("summary_sha256") != DECODED_TRANSFORM_SUMMARY_SHA256
        or source.get("manifest_sha256") != DECODED_TRANSFORM_MANIFEST_SHA256
        or source.get("spec_sha256") != DECODED_TRANSFORM_SPEC_SHA256
        or source.get("contract_sha256") != DECODED_TRANSFORM_CONTRACT_SHA256
        or acceptance.get("fixed_patches") != FIXED25_COUNT
        or acceptance.get("wsi_clusters") != 16
        or acceptance.get("all_exact_d4_algebra_checks") is not True
    ):
        message = "El addendum de Spec 0051 no corresponde al resultado aceptado"
        raise ValueError(message)
    if publication.get("paper_updated") is not False or publication.get(
        "sealed_test_updated",
        False,
    ):
        message = "El estado de publicación de Spec 0051 es incompatible"
        raise ValueError(message)

    dense_review = addendum.get("dense_rotation")
    if not isinstance(dense_review, dict):
        message = "El addendum no contiene el resumen denso de Spec 0051"
        raise TypeError(message)
    review_metrics = {
        "action_ratio_median": "action_ratio",
        "canonical_ratio_median": "canonical_ratio",
    }
    for review_name, metric_name in review_metrics.items():
        review_metric = dense_review.get(review_name)
        metric = comparisons[metric_name]
        paired = metric["paired"]
        if not isinstance(review_metric, dict):
            message = f"Falta el resumen revisado {review_name}"
            raise TypeError(message)
        expected_values = {
            "normal": float(metric["normal_median"]),
            "so2": float(metric["so2_median"]),
            "relative_so2_reduction": 1.0
            - float(metric["so2_median"]) / float(metric["normal_median"]),
            "wsi_cluster_median_difference_so2_minus_normal": float(
                paired["cluster_median_difference_so2_minus_normal"],
            ),
        }
        for name, expected in expected_values.items():
            observed = review_metric.get(name)
            if not isinstance(observed, (int, float)) or not math.isclose(
                float(observed),
                expected,
                rel_tol=0.0,
                abs_tol=1e-8,
            ):
                message = f"Resumen revisado inconsistente: {review_name}.{name}"
                raise ValueError(message)
        if (
            review_metric.get("so2_favorable_patches")
            != paired["so2_favorable_patch_count"]
            or review_metric.get("so2_favorable_wsi_clusters")
            != paired["so2_favorable_cluster_count"]
            or review_metric.get("descriptive_interval")
            != [paired["interval_low"], paired["interval_high"]]
        ):
            message = f"Conteos o intervalo inconsistentes: {review_name}"
            raise ValueError(message)

    input_review = dense_review.get("end_to_end_input_commutation_ratio_median")
    input_metric = comparisons["input_commutation_ratio"]
    input_paired = input_metric["paired"]
    if not isinstance(input_review, dict):
        message = "Falta el resumen revisado de conmutación entrada--salida"
        raise TypeError(message)
    input_expected = {
        "normal": input_metric["normal_median"],
        "so2": input_metric["so2_median"],
        "so2_favorable_patches": input_paired["so2_favorable_patch_count"],
        "wsi_cluster_median_difference_so2_minus_normal": input_paired[
            "cluster_median_difference_so2_minus_normal"
        ],
        "descriptive_interval": [
            input_paired["interval_low"],
            input_paired["interval_high"],
        ],
    }
    if any(
        input_review.get(name) != expected for name, expected in input_expected.items()
    ):
        message = "El resumen revisado de entrada--salida es inconsistente"
        raise ValueError(message)

    quarter_review = addendum.get("exact_quarter_rotations")
    if not isinstance(quarter_review, dict):
        message = "Falta el resumen revisado de rotaciones C4 exactas"
        raise TypeError(message)
    if quarter_review.get("aggregate_action_ratio_median") != {
        "normal": observed_quarter["normal_action"],
        "so2": observed_quarter["so2_action"],
    } or quarter_review.get("aggregate_canonical_ratio_median") != {
        "normal": observed_quarter["normal_canonical"],
        "so2": observed_quarter["so2_canonical"],
    }:
        message = "El addendum C4 no coincide con el resumen aceptado"
        raise ValueError(message)
    return document


def font(size: int, *, bold: bool = False) -> ImageFont.FreeTypeFont:
    names = [
        "/usr/share/fonts/truetype/liberation2/LiberationSans-Bold.ttf"
        if bold
        else "/usr/share/fonts/truetype/liberation2/LiberationSans-Regular.ttf",
        "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"
        if bold
        else "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
    ]
    for name in names:
        if Path(name).exists():
            return ImageFont.truetype(name, size=size)
    return ImageFont.load_default()


def es_decimal(value: float, digits: int) -> str:
    """Format one decimal value for Spanish prose."""
    return f"{value:.{digits}f}".replace(".", ",")


def draw_dashed_horizontal(
    draw: ImageDraw.ImageDraw,
    left: int,
    right: int,
    y: int,
    *,
    fill: str,
) -> None:
    """Draw one horizontal dashed reference line."""
    cursor = left
    while cursor < right:
        draw.line((cursor, y, min(cursor + 16, right), y), fill=fill, width=3)
        cursor += 26


def draw_orbit_population_report_sheets(
    source: Path,
    first_output: Path,
    second_output: Path,
) -> None:
    """Reflow the accepted 25-patch orbit grid into two readable report sheets.

    Raises:
        ValueError: The immutable source image has an unexpected size.

    """
    with Image.open(source) as opened:
        source_image = opened.convert("RGB")
    if source_image.size != (4320, 2340):
        message = f"Tamaño inesperado para el mosaico 25x360: {source_image.size}"
        raise ValueError(message)

    def render(indices: range, output: Path, subtitle: str) -> None:
        canvas = Image.new("RGB", (1800, 1680), "white")
        draw = ImageDraw.Draw(canvas)
        draw.text(
            (70, 24),
            "Órbitas latentes de los 25 parches fijos",
            fill="#1F2933",
            font=font(38, bold=True),
        )
        draw.text((70, 76), subtitle, fill="#52606D", font=font(25))
        is_first_sheet = len(indices) == FIRST_ORBIT_SHEET_COUNT
        columns = 3 if is_first_sheet else 2
        x_positions = [70, 640, 1210] if is_first_sheet else [350, 930]
        label_font = font(30, bold=True)
        for position, patch_index in enumerate(indices):
            source_row, source_column = divmod(patch_index, 5)
            source_left = source_column * 864
            source_top = 60 + source_row * 456
            tile = source_image.crop(
                (source_left, source_top, source_left + 864, source_top + 456),
            )
            tile = tile.resize((520, 274), Image.Resampling.LANCZOS)
            tile_draw = ImageDraw.Draw(tile)
            tile_draw.rectangle((0, 0, 520, 40), fill="white")
            tile_draw.text(
                (8, 3),
                f"P{patch_index:02d} · normal",
                fill=f"#{BLUE}",
                font=label_font,
            )
            tile_draw.text(
                (278, 3),
                f"P{patch_index:02d} · SO(2)",
                fill=f"#{ORANGE}",
                font=label_font,
            )
            destination_column = position % columns
            destination_row = position // columns
            canvas.paste(
                tile,
                (x_positions[destination_column], 130 + destination_row * 300),
            )
        draw.text(
            (70, 1640),
            "Cada curva contiene 360 posiciones (0°--359°); la escala es isótropa y propia de cada panel.",
            fill="#52606D",
            font=font(27),
        )
        canvas.save(output, dpi=(180, 180))

    render(range(0, 15), first_output, "Parte A · parches 00--14 · muestreo cada 1°")
    render(range(15, 25), second_output, "Parte B · parches 15--24 · muestreo cada 1°")


def draw_population_controls_report_figure(source: Path, output: Path) -> None:
    """Replace source-population shorthand in a report-only derived figure.

    Raises:
        ValueError: If the immutable source artifact has an unexpected size.

    """
    with Image.open(source) as opened:
        image = opened.convert("RGB")
    if image.size != (3420, 900):
        message = f"Tamaño inesperado para los controles rotacionales: {image.size}"
        raise ValueError(message)
    draw = ImageDraw.Draw(image)
    title = "Métricas corregidas de los 25 parches fijos"
    title_font = font(42)
    bounds = draw.textbbox((0, 0), title, font=title_font)
    title_width = bounds[2] - bounds[0]
    left = (image.width - title_width) // 2
    draw.rectangle((left - 25, 0, left + title_width + 25, 48), fill="white")
    draw.text((left, 0), title, fill="black", font=title_font)
    image.save(output)


def draw_decoded_transform_population_figure(
    document: dict[str, Any],
    output: Path,
) -> None:
    """Draw untruncated paired dense-ratio comparisons from Spec 0051."""
    image = Image.new("RGB", (1800, 880), "white")
    draw = ImageDraw.Draw(image)
    title_font = font(38, bold=True)
    subtitle_font = font(24)
    axis_font = font(22)
    small_font = font(19)
    draw.text(
        (70, 28),
        "Consistencia decodificada en los 25 parches fijos",
        fill=f"#{INK}",
        font=title_font,
    )
    draw.text(
        (70, 80),
        "Comparación pareada normal vs. SO(2) · rotaciones 0°, 5°, ..., 355°",
        fill=f"#{MUTED}",
        font=subtitle_font,
    )
    panels = [
        ("action_ratio", "Acción espacial del embedding"),
        ("canonical_ratio", "Desrotación del embedding"),
    ]
    for panel_index, (metric_name, panel_title) in enumerate(panels):
        panel_left = 90 + panel_index * 850
        panel_right = panel_left + 760
        plot_top, plot_bottom = 175, 690
        x_normal, x_so2 = panel_left + 245, panel_left + 555
        draw.text(
            (panel_left + 25, 128),
            panel_title,
            fill=f"#{INK}",
            font=font(27, bold=True),
        )
        for tick in (0.0, 0.25, 0.5, 0.75, 1.0):
            y = plot_bottom - int(tick * (plot_bottom - plot_top))
            if math.isclose(tick, 0.5):
                draw_dashed_horizontal(
                    draw,
                    panel_left + 95,
                    panel_right - 25,
                    y,
                    fill="#596675",
                )
            else:
                draw.line(
                    (panel_left + 95, y, panel_right - 25, y),
                    fill="#E3E8EE",
                    width=2,
                )
            draw.text(
                (panel_left + 28, y - 12),
                f"{tick:.2f}".replace(".", ","),
                fill=f"#{MUTED}",
                font=small_font,
            )
        metric = document["comparisons"][metric_name]
        normal_values = [float(value) for value in metric["normal"]]
        so2_values = [float(value) for value in metric["so2"]]
        for normal_value, so2_value in zip(normal_values, so2_values, strict=True):
            y_normal = plot_bottom - int(normal_value * (plot_bottom - plot_top))
            y_so2 = plot_bottom - int(so2_value * (plot_bottom - plot_top))
            draw.line(
                (x_normal, y_normal, x_so2, y_so2),
                fill="#AAB5C1",
                width=2,
            )
            draw.ellipse(
                (x_normal - 5, y_normal - 5, x_normal + 5, y_normal + 5),
                fill=f"#{BLUE}",
            )
            draw.ellipse(
                (x_so2 - 5, y_so2 - 5, x_so2 + 5, y_so2 + 5),
                fill=f"#{ORANGE}",
            )
        for x_value, branch in (
            (x_normal, "normal"),
            (x_so2, "so2"),
        ):
            value = float(metric[f"{branch}_median"])
            y = plot_bottom - int(value * (plot_bottom - plot_top))
            draw.line((x_value - 34, y, x_value + 34, y), fill="#111827", width=7)
            label = f"mediana {value:.3f}".replace(".", ",")
            bounds = draw.textbbox((0, 0), label, font=small_font)
            draw.text(
                (x_value - (bounds[2] - bounds[0]) / 2, y - 34),
                label,
                fill="#111827",
                font=small_font,
            )
        draw.text(
            (x_normal - 88, plot_bottom + 24),
            "Normal",
            fill=f"#{BLUE}",
            font=axis_font,
        )
        draw.text(
            (x_so2 - 50, plot_bottom + 24),
            "SO(2)",
            fill=f"#{ORANGE}",
            font=axis_font,
        )
        draw.text(
            (panel_left + 235, 755),
            f"SO(2) menor en {metric['paired']['so2_favorable_patch_count']}/{metric['paired']['n_patches']}",
            fill=f"#{INK}",
            font=font(22, bold=True),
        )
    draw.text(
        (70, 830),
        "Razón 1,00 = no actuar · línea gris discontinua: criterio absoluto 0,50 = conservar como máximo la mitad del error.",
        fill=f"#{MUTED}",
        font=small_font,
    )
    image.save(output, dpi=(180, 180))


def draw_decoded_transform_exact_figure(
    document: dict[str, Any],
    output: Path,
) -> None:
    """Draw exact C4 and reflection medians with the absolute reference."""
    image = Image.new("RGB", (1800, 920), "white")
    draw = ImageDraw.Draw(image)
    draw.text(
        (65, 25),
        "Transformaciones exactas en la cuadrícula",
        fill=f"#{INK}",
        font=font(38, bold=True),
    )
    draw.text(
        (65, 77),
        "Medianas sobre 25 parches · C4 y cuatro reflexiones",
        fill=f"#{MUTED}",
        font=font(24),
    )
    transforms = [
        ("rot90", "R90"),
        ("rot180", "R180"),
        ("rot270", "R270"),
        ("flip_h", "Fh"),
        ("flip_v", "Fv"),
        ("flip_diag", "Fd"),
        ("flip_anti_diag", "Fa"),
    ]
    panels = [
        ("action_ratio_median_valid", "Acción latente"),
        ("canonical_ratio_median_valid", "Canonicalización"),
        ("input_commutation_ratio_median_valid", "Entrada a salida"),
    ]
    for panel_index, (metric_name, title) in enumerate(panels):
        panel_left = 45 + panel_index * 585
        plot_left, plot_right = panel_left + 68, panel_left + 560
        plot_top, plot_bottom = 185, 710
        draw.text(
            (panel_left + 105, 130),
            title,
            fill=f"#{INK}",
            font=font(25, bold=True),
        )
        for tick in (0.0, 0.25, 0.5, 0.75):
            y = plot_bottom - int(tick / 0.8 * (plot_bottom - plot_top))
            if math.isclose(tick, 0.5):
                draw_dashed_horizontal(
                    draw,
                    plot_left,
                    plot_right,
                    y,
                    fill="#596675",
                )
            else:
                draw.line(
                    (plot_left, y, plot_right, y),
                    fill="#E3E8EE",
                    width=2,
                )
            draw.text(
                (panel_left + 4, y - 11),
                f"{tick:.2f}".replace(".", ","),
                fill=f"#{MUTED}",
                font=font(17),
            )
        group_width = (plot_right - plot_left) / len(transforms)
        for transform_index, (transform, short_label) in enumerate(transforms):
            center = plot_left + int((transform_index + 0.5) * group_width)
            for branch, x_offset, color in (
                ("normal_vae", -16, BLUE),
                ("so2_vae", 7, ORANGE),
            ):
                value = float(
                    document["branches"][branch]["exact_d4"][transform][metric_name],
                )
                bar_top = plot_bottom - int(
                    min(value, 0.8) / 0.8 * (plot_bottom - plot_top),
                )
                draw.rectangle(
                    (center + x_offset, bar_top, center + x_offset + 18, plot_bottom),
                    fill=f"#{color}",
                )
            bounds = draw.textbbox((0, 0), short_label, font=font(17))
            draw.text(
                (center - (bounds[2] - bounds[0]) / 2, plot_bottom + 18),
                short_label,
                fill=f"#{INK}",
                font=font(17),
            )
    draw.rectangle((545, 810, 585, 834), fill=f"#{BLUE}")
    draw.text((598, 806), "VAE normal", fill=f"#{INK}", font=font(21))
    draw.rectangle((895, 810, 935, 834), fill=f"#{ORANGE}")
    draw.text((948, 806), "VAE SO(2)", fill=f"#{INK}", font=font(21))
    draw.text(
        (340, 865),
        "Fh/Fv: reflexión horizontal/vertical · Fd/Fa: diagonal principal/antidiagonal · línea gris discontinua: criterio absoluto 0,50",
        fill=f"#{MUTED}",
        font=font(19),
    )
    image.save(output, dpi=(180, 180))


def draw_mil_chart(normal: dict[str, Any], so2: dict[str, Any], output: Path) -> None:
    width, height = 1500, 760
    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)
    title_font, axis_font, tick_font = font(38, bold=True), font(25), font(22)
    draw.text(
        (80, 38),
        es("Diagnostico WSI en el test sellado"),
        fill="#1F2933",
        font=title_font,
    )
    left, top, right, bottom = 130, 150, 1430, 610
    draw.line((left, top, left, bottom), fill="#8392A5", width=3)
    draw.line((left, bottom, right, bottom), fill="#8392A5", width=3)
    for i in range(0, 8):
        value = i / 10
        y = bottom - int(value / 0.7 * (bottom - top))
        draw.line((left, y, right, y), fill="#E5E9EF", width=2)
        draw.text((54, y - 13), f"{value:.1f}", fill="#52606D", font=tick_font)
    metrics = [
        ("F1 macro", normal["metrics"]["macro_f1"], so2["metrics"]["macro_f1"]),
        (
            "Exactitud\nbalanceada",
            normal["metrics"]["balanced_accuracy"],
            so2["metrics"]["balanced_accuracy"],
        ),
        ("Exactitud", normal["metrics"]["accuracy"], so2["metrics"]["accuracy"]),
    ]
    centers = [360, 790, 1220]
    bar_width = 110
    for center, (label, nval, sval) in zip(centers, metrics, strict=True):
        for x, val, color in ((center - 125, nval, BLUE), (center + 20, sval, ORANGE)):
            y = bottom - int(val / 0.7 * (bottom - top))
            draw.rounded_rectangle(
                (x, y, x + bar_width, bottom),
                radius=8,
                fill=f"#{color}",
            )
            text = f"{val:.3f}"
            bbox = draw.textbbox((0, 0), text, font=tick_font)
            draw.text(
                (x + (bar_width - (bbox[2] - bbox[0])) / 2, y - 34),
                text,
                fill="#1F2933",
                font=tick_font,
            )
        for line_no, line in enumerate(label.split("\n")):
            bbox = draw.textbbox((0, 0), line, font=axis_font)
            draw.text(
                (center - (bbox[2] - bbox[0]) / 2, bottom + 22 + line_no * 28),
                es(line),
                fill="#1F2933",
                font=axis_font,
            )
    legend_y = 700
    draw.rounded_rectangle(
        (395, legend_y, 435, legend_y + 24),
        radius=4,
        fill=f"#{BLUE}",
    )
    draw.text((450, legend_y - 4), NORMAL_LABEL, fill="#1F2933", font=tick_font)
    draw.rounded_rectangle(
        (850, legend_y, 890, legend_y + 24),
        radius=4,
        fill=f"#{ORANGE}",
    )
    draw.text((905, legend_y - 4), SO2_LABEL, fill="#1F2933", font=tick_font)
    image.save(output, dpi=(180, 180))


def draw_tissue_chart(table: dict[str, Any], output: Path) -> None:
    width, height = 1500, 800
    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)
    title_font, axis_font, tick_font = font(38, bold=True), font(25), font(22)
    draw.text(
        (80, 35),
        "Eficiencia de etiquetas en tejido: F1 macro",
        fill="#1F2933",
        font=title_font,
    )
    left, top, right, bottom = 135, 150, 1435, 635
    draw.line((left, top, left, bottom), fill="#8392A5", width=3)
    draw.line((left, bottom, right, bottom), fill="#8392A5", width=3)
    y_min, y_max = 0.45, 0.80
    for value in [0.45, 0.50, 0.55, 0.60, 0.65, 0.70, 0.75, 0.80]:
        y = bottom - int((value - y_min) / (y_max - y_min) * (bottom - top))
        draw.line((left, y, right, y), fill="#E5E9EF", width=2)
        draw.text((55, y - 13), f"{value:.2f}", fill="#52606D", font=tick_font)
    rows = table["rows"]
    xs = [
        left + 95 + i * ((right - left - 190) / (len(rows) - 1))
        for i in range(len(rows))
    ]

    def y_coord(value: float) -> float:
        return bottom - (value - y_min) / (y_max - y_min) * (bottom - top)

    for key, color in (("normal_vae", BLUE), ("so2_vae", ORANGE)):
        points = [
            (x, y_coord(row[key]["metrics"]["macro_f1"]))
            for x, row in zip(xs, rows, strict=True)
        ]
        draw.line(points, fill=f"#{color}", width=7, joint="curve")
        for x, y in points:
            marker = (x - 10, y - 10, x + 10, y + 10)
            if key == "normal_vae":
                draw.ellipse(marker, fill=f"#{color}", outline="white", width=3)
            else:
                draw.rectangle(marker, fill=f"#{color}", outline="white", width=3)
    for x, row in zip(xs, rows, strict=True):
        label = f"{row['labels_per_class']:,}".replace(",", ".")
        bbox = draw.textbbox((0, 0), label, font=tick_font)
        draw.text(
            (x - (bbox[2] - bbox[0]) / 2, bottom + 20),
            label,
            fill="#1F2933",
            font=tick_font,
        )
    sig = rows[1]
    sx = xs[1]
    sy = (
        min(
            y_coord(sig["normal_vae"]["metrics"]["macro_f1"]),
            y_coord(sig["so2_vae"]["metrics"]["macro_f1"]),
        )
        - 38
    )
    draw.text((sx - 10, sy), "*", fill="#7C3AED", font=font(38, bold=True))
    draw.text((left + 455, 690), "Etiquetas por clase", fill="#1F2933", font=axis_font)
    draw.rounded_rectangle((845, 683, 885, 707), radius=4, fill=f"#{BLUE}")
    draw.text((900, 679), NORMAL_LABEL, fill="#1F2933", font=tick_font)
    draw.rounded_rectangle((845, 728, 885, 752), radius=4, fill=f"#{ORANGE}")
    draw.text((900, 724), SO2_LABEL, fill="#1F2933", font=tick_font)
    draw.text(
        (150, 715),
        es("* intervalo simultaneo del contraste excluye cero"),
        fill="#7C3AED",
        font=tick_font,
    )
    image.save(output, dpi=(180, 180))


def draw_wsi_class_chart(
    normal: dict[str, Any],
    so2: dict[str, Any],
    output: Path,
) -> None:
    width, height = 1800, 1450
    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)
    title_font = font(40, bold=True)
    subtitle_font = font(25)
    axis_font = font(24, bold=True)
    tick_font = font(21)
    cell_font = font(19, bold=True)
    class_labels = [item["diagnosis"] for item in normal["metrics"]["per_class"]]
    supports = [item["support"] for item in normal["metrics"]["per_class"]]

    draw.text(
        (70, 35),
        "Diagnóstico WSI: errores y F1 por clase",
        fill="#1F2933",
        font=title_font,
    )
    draw.text(
        (70, 90),
        "Test sellado: n=23 WSI; las celdas muestran conteo y porcentaje dentro de la clase real",
        fill="#52606D",
        font=subtitle_font,
    )

    def heat_color(value: float) -> str:
        start = (239, 246, 251)
        end = (46, 111, 158)
        rgb = tuple(round(a + value * (b - a)) for a, b in zip(start, end, strict=True))
        return "#" + "".join(f"{channel:02X}" for channel in rgb)

    def draw_confusion(metrics: dict[str, Any], x0: int, title: str) -> None:
        matrix = metrics["confusion_matrix"]
        cell = 92
        y0 = 260
        draw.text((x0 + 95, 150), title, fill="#1F2933", font=axis_font)
        draw.text((x0 + 150, 190), "Predicción", fill="#52606D", font=tick_font)
        for col, label in enumerate(class_labels):
            bbox = draw.textbbox((0, 0), label, font=tick_font)
            draw.text(
                (x0 + col * cell + (cell - (bbox[2] - bbox[0])) / 2, y0 - 34),
                label,
                fill="#1F2933",
                font=tick_font,
            )
        for row, (values, support) in enumerate(zip(matrix, supports, strict=True)):
            label = f"{class_labels[row]} (n={support})"
            bbox = draw.textbbox((0, 0), label, font=tick_font)
            draw.text(
                (x0 - (bbox[2] - bbox[0]) - 14, y0 + row * cell + 32),
                label,
                fill="#1F2933",
                font=tick_font,
            )
            for col, count in enumerate(values):
                share = count / support if support else 0.0
                left = x0 + col * cell
                top = y0 + row * cell
                draw.rectangle(
                    (left, top, left + cell, top + cell),
                    fill=heat_color(share),
                    outline="#FFFFFF",
                    width=3,
                )
                label = f"{count}\n{share:.0%}"
                color = "white" if share >= HEATMAP_LIGHT_TEXT_THRESHOLD else "#1F2933"
                for line_no, line in enumerate(label.split("\n")):
                    bbox = draw.textbbox((0, 0), line, font=cell_font)
                    draw.text(
                        (
                            left + (cell - (bbox[2] - bbox[0])) / 2,
                            top + 20 + line_no * 25,
                        ),
                        line,
                        fill=color,
                        font=cell_font,
                    )

    draw_confusion(normal["metrics"], 235, NORMAL_LABEL)
    draw_confusion(so2["metrics"], 1120, SO2_LABEL)

    left, top, right, bottom = 120, 880, 1710, 1280
    draw.text((70, 785), "F1 por clase", fill="#1F2933", font=axis_font)
    draw.line((left, top, left, bottom), fill="#8392A5", width=3)
    draw.line((left, bottom, right, bottom), fill="#8392A5", width=3)
    for value in [0.0, 0.25, 0.5, 0.75, 1.0]:
        y = bottom - int(value * (bottom - top))
        draw.line((left, y, right, y), fill="#E5E9EF", width=2)
        draw.text((55, y - 12), f"{value:.2f}", fill="#52606D", font=tick_font)
    normal_f1 = [item["f1"] for item in normal["metrics"]["per_class"]]
    so2_f1 = [item["f1"] for item in so2["metrics"]["per_class"]]
    group_width = (right - left) / len(class_labels)
    bar_width = 92
    for index, (label, support, nval, sval) in enumerate(
        zip(class_labels, supports, normal_f1, so2_f1, strict=True),
    ):
        center = left + group_width * (index + 0.5)
        for x, value, color in (
            (center - bar_width - 8, nval, BLUE),
            (center + 8, sval, ORANGE),
        ):
            y = bottom - value * (bottom - top)
            draw.rounded_rectangle(
                (x, y, x + bar_width, bottom),
                radius=7,
                fill=f"#{color}",
            )
            text = f"{value:.2f}"
            bbox = draw.textbbox((0, 0), text, font=tick_font)
            draw.text(
                (x + (bar_width - (bbox[2] - bbox[0])) / 2, y - 30),
                text,
                fill="#1F2933",
                font=tick_font,
            )
        group_label = f"{label}\n(n={support})"
        for line_no, line in enumerate(group_label.split("\n")):
            bbox = draw.textbbox((0, 0), line, font=tick_font)
            draw.text(
                (center - (bbox[2] - bbox[0]) / 2, bottom + 18 + line_no * 27),
                line,
                fill="#1F2933",
                font=tick_font,
            )
    draw.rounded_rectangle((520, 1380, 560, 1404), radius=4, fill=f"#{BLUE}")
    draw.text((575, 1376), NORMAL_LABEL, fill="#1F2933", font=tick_font)
    draw.rounded_rectangle((1040, 1380, 1080, 1404), radius=4, fill=f"#{ORANGE}")
    draw.text((1095, 1376), SO2_LABEL, fill="#1F2933", font=tick_font)
    image.save(output, dpi=(180, 180))


def draw_tissue_class_chart(table: dict[str, Any], output: Path) -> None:
    width, height = 1800, 900
    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)
    title_font = font(40, bold=True)
    axis_font = font(23, bold=True)
    tick_font = font(19)
    draw.text(
        (70, 35),
        "Eficiencia de etiquetas: F1 por tipo de tejido",
        fill="#1F2933",
        font=title_font,
    )
    draw.text(
        (70, 90),
        "Test sellado: 31.572 parches de 23 WSI; puntos descriptivos por presupuesto",
        fill="#52606D",
        font=font(24),
    )
    rows = table["rows"]
    tissues = [
        ("tumor", "Tumor", "n=21.796 parches; 23 WSI"),
        ("stroma", "Estroma", "n=8.500 parches; 21 WSI"),
        ("necrosis", "Necrosis", "n=1.276 parches; 5 WSI"),
    ]

    for panel, (key, label, support) in enumerate(tissues):
        panel_left = 75 + panel * 575
        left, top, right, bottom = panel_left + 60, 240, panel_left + 515, 700
        draw.text((panel_left + 150, 160), label, fill="#1F2933", font=axis_font)
        draw.text((panel_left + 110, 195), support, fill="#52606D", font=tick_font)
        draw.line((left, top, left, bottom), fill="#8392A5", width=3)
        draw.line((left, bottom, right, bottom), fill="#8392A5", width=3)
        for value in [0.0, 0.25, 0.5, 0.75, 1.0]:
            y = bottom - int(value * (bottom - top))
            draw.line((left, y, right, y), fill="#E5E9EF", width=2)
            draw.text(
                (left - 48, y - 10),
                f"{value:.2f}",
                fill="#52606D",
                font=tick_font,
            )
        xs = [left + i * (right - left) / (len(rows) - 1) for i in range(len(rows))]
        for branch, color in (("normal_vae", BLUE), ("so2_vae", ORANGE)):
            values = []
            for row in rows:
                by_tissue = {
                    item["tissue"]: item["f1"]
                    for item in row[branch]["metrics"]["per_class"]
                }
                values.append(by_tissue[key])
            points = [
                (x, bottom - value * (bottom - top))
                for x, value in zip(xs, values, strict=True)
            ]
            draw.line(points, fill=f"#{color}", width=6, joint="curve")
            for x, y in points:
                marker = (x - 8, y - 8, x + 8, y + 8)
                if branch == "normal_vae":
                    draw.ellipse(marker, fill=f"#{color}", outline="white", width=2)
                else:
                    draw.rectangle(marker, fill=f"#{color}", outline="white", width=2)
        for x, row in zip(xs, rows, strict=True):
            value = f"{row['labels_per_class']:,}".replace(",", ".")
            bbox = draw.textbbox((0, 0), value, font=tick_font)
            draw.text(
                (x - (bbox[2] - bbox[0]) / 2, bottom + 18),
                value,
                fill="#1F2933",
                font=tick_font,
            )
    draw.text((710, 765), "Etiquetas por clase", fill="#1F2933", font=axis_font)
    draw.rounded_rectangle((490, 835, 530, 859), radius=4, fill=f"#{BLUE}")
    draw.text((545, 831), NORMAL_LABEL, fill="#1F2933", font=tick_font)
    draw.rounded_rectangle((1030, 835, 1070, 859), radius=4, fill=f"#{ORANGE}")
    draw.text((1085, 831), SO2_LABEL, fill="#1F2933", font=tick_font)
    image.save(output, dpi=(180, 180))


def draw_supervised_development_chart(
    mil_train: dict[str, list[dict[str, str]]],
    mil_validation: dict[str, list[dict[str, str]]],
    mil_summaries: dict[str, dict[str, Any]],
    tissue_validation: list[dict[str, float]],
    output: Path,
) -> None:
    width, height = 2200, 1200
    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)
    title_font = font(46, bold=True)
    axis_font = font(30, bold=True)
    tick_font = font(26)
    draw.text(
        (70, 35),
        "Desarrollo de los clasificadores supervisados",
        fill="#1F2933",
        font=title_font,
    )
    draw.text(
        (70, 90),
        "F1 de selección y brecha train-validación; el test sellado permanece separado",
        fill="#52606D",
        font=font(28),
    )

    def y_at(
        value: float,
        top: int,
        bottom: int,
        minimum: float,
        maximum: float,
    ) -> float:
        return bottom - (value - minimum) / (maximum - minimum) * (bottom - top)

    def x_at(boundary: float, left: int, right: int, maximum: float) -> float:
        return left + boundary / maximum * (right - left)

    def axes(  # noqa: PLR0913, PLR0917
        left: int,
        top: int,
        right: int,
        bottom: int,
        title: str,
        y_values: list[float],
    ) -> None:
        draw.text((left, top - 62), title, fill="#1F2933", font=axis_font)
        draw.line((left, top, left, bottom), fill="#8392A5", width=3)
        draw.line((left, bottom, right, bottom), fill="#8392A5", width=3)
        y_min, y_max = min(y_values), max(y_values)
        for value in y_values:
            y = y_at(value, top, bottom, y_min, y_max)
            draw.line((left, y, right, y), fill="#E5E9EF", width=2)
            draw.text(
                (left - 58, y - 10),
                f"{value:.2f}",
                fill="#52606D",
                font=tick_font,
            )

    def dashed_line(
        points: list[tuple[float, float]],
        color: str,
        width: int = 4,
        dash: int = 18,
        gap: int = 10,
    ) -> None:
        for (x1, y1), (x2, y2) in pairwise(points):
            distance = max(((x2 - x1) ** 2 + (y2 - y1) ** 2) ** 0.5, 1.0)
            cursor = 0.0
            while cursor < distance:
                segment_end = min(cursor + dash, distance)
                start_fraction = cursor / distance
                end_fraction = segment_end / distance
                draw.line(
                    (
                        x1 + (x2 - x1) * start_fraction,
                        y1 + (y2 - y1) * start_fraction,
                        x1 + (x2 - x1) * end_fraction,
                        y1 + (y2 - y1) * end_fraction,
                    ),
                    fill=color,
                    width=width,
                )
                cursor += dash + gap

    max_boundary = max(
        float(row["boundary"]) for rows in mil_validation.values() for row in rows
    )
    left, top, right, bottom = 105, 230, 1030, 630
    axes(
        left,
        top,
        right,
        bottom,
        "MIL WSI: F1 macro de validación",
        [0.0, 0.25, 0.5, 0.75, 1.0],
    )
    for branch, color in (("normal_vae", BLUE), ("so2_vae", ORANGE)):
        points = [
            (
                x_at(float(row["boundary"]), left, right, max_boundary),
                y_at(float(row["macro_f1"]), top, bottom, 0.0, 1.0),
            )
            for row in mil_validation[branch]
        ]
        draw.line(points, fill=f"#{color}", width=4, joint="curve")
        summary = mil_summaries[branch]
        best_x = x_at(summary["best_boundary"], left, right, max_boundary)
        best_y = y_at(
            summary["best_metrics"]["macro_f1"],
            top,
            bottom,
            0.0,
            1.0,
        )
        draw.ellipse(
            (best_x - 11, best_y - 11, best_x + 11, best_y + 11),
            fill=f"#{color}",
            outline="white",
            width=3,
        )
        label = f"{summary['best_metrics']['macro_f1']:.3f} @ {summary['best_boundary']:,}".replace(
            ",",
            ".",
        )
        draw.text((best_x + 12, best_y - 12), label, fill=f"#{color}", font=tick_font)
    for value in [0, 2000, 4000, 6000, 8000]:
        x = x_at(value, left, right, max_boundary)
        draw.text(
            (x - 25, bottom + 18),
            f"{value // 1000}k",
            fill="#52606D",
            font=tick_font,
        )
    draw.text(
        (left + 365, bottom + 58),
        "Actualización",
        fill="#52606D",
        font=tick_font,
    )

    left2, top2, right2, bottom2 = 1190, 230, 2115, 630
    axes(
        left2,
        top2,
        right2,
        bottom2,
        "MIL WSI: entropía cruzada train-validación",
        [0.5, 1.0, 1.5, 2.0, 2.5],
    )
    for branch, color in (("normal_vae", BLUE), ("so2_vae", ORANGE)):
        previous_boundary = 0
        train_points: list[tuple[float, float]] = []
        validation_points: list[tuple[float, float]] = []
        for validation_row in mil_validation[branch]:
            boundary = int(validation_row["boundary"])
            losses = [
                float(row["unweighted_loss"])
                for row in mil_train[branch]
                if previous_boundary < int(row["committed_update"]) <= boundary
            ]
            if not losses:
                message = f"Missing MIL train-loss window ending at {boundary}"
                raise ValueError(message)
            train_points.append(
                (
                    x_at(boundary, left2, right2, max_boundary),
                    y_at(sum(losses) / len(losses), top2, bottom2, 0.5, 2.5),
                ),
            )
            validation_points.append(
                (
                    x_at(boundary, left2, right2, max_boundary),
                    y_at(
                        float(validation_row["mean_ce"]),
                        top2,
                        bottom2,
                        0.5,
                        2.5,
                    ),
                ),
            )
            previous_boundary = boundary
        draw.line(train_points, fill=f"#{color}", width=5, joint="curve")
        dashed_line(validation_points, f"#{color}")
    for value in [0, 2000, 4000, 6000, 8000]:
        x = x_at(value, left2, right2, max_boundary)
        draw.text(
            (x - 25, bottom2 + 18),
            f"{value // 1000}k",
            fill="#52606D",
            font=tick_font,
        )
    draw.text(
        (left2 + 365, bottom2 + 58),
        "Actualización",
        fill="#52606D",
        font=tick_font,
    )

    draw.rounded_rectangle((500, 730, 540, 754), radius=4, fill=f"#{BLUE}")
    draw.text((555, 726), NORMAL_LABEL, fill="#1F2933", font=tick_font)
    draw.rounded_rectangle((1090, 730, 1130, 754), radius=4, fill=f"#{ORANGE}")
    draw.text((1145, 726), SO2_LABEL, fill="#1F2933", font=tick_font)
    draw.line((590, 790, 650, 790), fill="#52606D", width=5)
    draw.text(
        (665, 780),
        "train online (media por ventana)",
        fill="#1F2933",
        font=tick_font,
    )
    dashed_line([(1260, 790), (1320, 790)], "#52606D", width=4)
    draw.text((1335, 780), "validación", fill="#1F2933", font=tick_font)

    left3, top3, right3, bottom3 = 420, 920, 1780, 1060
    axes(
        left3,
        top3,
        right3,
        bottom3,
        "Tejido: mejor F1 macro de validación",
        [0.0, 0.5, 1.0],
    )
    xs = [
        left3 + i * (right3 - left3) / (len(tissue_validation) - 1)
        for i in range(len(tissue_validation))
    ]
    for key, color in (("normal_vae", BLUE), ("so2_vae", ORANGE)):
        points = [
            (x, y_at(row[key], top3, bottom3, 0.0, 1.0))
            for x, row in zip(xs, tissue_validation, strict=True)
        ]
        draw.line(points, fill=f"#{color}", width=6, joint="curve")
        for x, y in points:
            if key == "normal_vae":
                draw.ellipse(
                    (x - 8, y - 8, x + 8, y + 8),
                    fill=f"#{color}",
                    outline="white",
                    width=2,
                )
            else:
                draw.rectangle(
                    (x - 8, y - 8, x + 8, y + 8),
                    fill=f"#{color}",
                    outline="white",
                    width=2,
                )
    for x, row in zip(xs, tissue_validation, strict=True):
        label = f"{int(row['budget']):,}".replace(",", ".")
        bbox = draw.textbbox((0, 0), label, font=tick_font)
        draw.text(
            (x - (bbox[2] - bbox[0]) / 2, bottom3 + 18),
            label,
            fill="#52606D",
            font=tick_font,
        )
    draw.text(
        (left3 + 540, bottom3 + 58),
        "Etiquetas por clase",
        fill="#52606D",
        font=tick_font,
    )
    image.save(output, dpi=(180, 180))


def set_cell_shading(cell: Any, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell: Any, color: str = "CDD4DC", size: str = "6") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_borders = tc_pr.first_child_found_in("w:tcBorders")
    if tc_borders is None:
        tc_borders = OxmlElement("w:tcBorders")
        tc_pr.append(tc_borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = tc_borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            tc_borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def set_cell_margins(
    cell: Any,
    top: int = 80,
    start: int = 90,
    bottom: int = 80,
    end: int = 90,
) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for key, value in (
        ("top", top),
        ("start", start),
        ("bottom", bottom),
        ("end", end),
    ):
        element = tc_mar.find(qn(f"w:{key}"))
        if element is None:
            element = OxmlElement(f"w:{key}")
            tc_mar.append(element)
        element.set(qn("w:w"), str(value))
        element.set(qn("w:type"), "dxa")


def set_cell_width(cell: Any, inches: float) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(inches * 1440)))
    tc_w.set(qn("w:type"), "dxa")


def set_repeat_table_header(row: Any) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_row_cant_split(row: Any) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def set_paragraph_keep(
    paragraph: Any,
    *,
    keep_next: bool = False,
    keep_lines: bool = True,
) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    if keep_next:
        p_pr.append(OxmlElement("w:keepNext"))
    if keep_lines:
        p_pr.append(OxmlElement("w:keepLines"))


def add_page_number(paragraph: Any) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    fld_char = OxmlElement("w:fldChar")
    fld_char.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char, instr_text, fld_char2])


def set_picture_alt(paragraph: Any, description: str) -> None:
    doc_prs = paragraph._p.xpath(".//wp:docPr")
    if doc_prs:
        doc_prs[-1].set("descr", description)
        doc_prs[-1].set("title", description[:80])


def add_table(
    doc: Document,
    headers: list[str],
    rows: list[list[str]],
    *,
    widths: list[float] | None = None,
    font_size: float = 8.7,
) -> Any:
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.style = "Table Grid"
    header = table.rows[0]
    set_repeat_table_header(header)
    for i, label in enumerate(headers):
        cell = header.cells[i]
        set_cell_shading(cell, INK)
        set_cell_border(cell)
        set_cell_margins(cell)
        if widths:
            set_cell_width(cell, widths[i])
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(es(label))
        run.bold = True
        run.font.color.rgb = RGBColor(255, 255, 255)
        run.font.size = Pt(font_size)
    for row_no, values in enumerate(rows):
        row = table.add_row()
        set_row_cant_split(row)
        for i, value in enumerate(values):
            cell = row.cells[i]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_border(cell)
            set_cell_margins(cell)
            if widths:
                set_cell_width(cell, widths[i])
            if row_no % 2:
                set_cell_shading(cell, "F5F7F9")
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.alignment = (
                WD_ALIGN_PARAGRAPH.LEFT if i == 0 else WD_ALIGN_PARAGRAPH.CENTER
            )
            run = p.add_run(es(str(value)))
            run.font.size = Pt(font_size)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_body(doc: Document, text: str, *, bold_prefix: str | None = None) -> Any:
    p = doc.add_paragraph()
    p.style = doc.styles["Body Text"]
    if bold_prefix and text.startswith(bold_prefix):
        p.add_run(es(bold_prefix)).bold = True
        p.add_run(es(text[len(bold_prefix) :]))
    else:
        p.add_run(es(text))
    return p


def add_bullet(doc: Document, text: str, *, level: int = 0) -> Any:
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    p.add_run(es(text))
    return p


def add_figure(doc: Document, path: Path, width: float, caption: str, alt: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(3)
    set_paragraph_keep(p, keep_next=True)
    p.add_run().add_picture(str(path), width=Inches(width))
    set_picture_alt(p, es(alt))
    cap = doc.add_paragraph(es(caption), style="Caption")
    cap.alignment = WD_ALIGN_PARAGRAPH.LEFT
    set_paragraph_keep(cap)


def add_heading(doc: Document, text: str, level: int = 1) -> Any:
    p = doc.add_heading(es(text), level=level)
    set_paragraph_keep(p, keep_next=True)
    return p


def set_document_language(doc: Document, language: str = "es-CO") -> None:
    styles_element = doc.styles.element
    doc_defaults = styles_element.find(qn("w:docDefaults"))
    if doc_defaults is None:
        doc_defaults = OxmlElement("w:docDefaults")
        styles_element.insert(0, doc_defaults)
    rpr_default = doc_defaults.find(qn("w:rPrDefault"))
    if rpr_default is None:
        rpr_default = OxmlElement("w:rPrDefault")
        doc_defaults.append(rpr_default)
    default_rpr = rpr_default.find(qn("w:rPr"))
    if default_rpr is None:
        default_rpr = OxmlElement("w:rPr")
        rpr_default.append(default_rpr)
    default_lang = default_rpr.find(qn("w:lang"))
    if default_lang is None:
        default_lang = OxmlElement("w:lang")
        default_rpr.append(default_lang)
    default_lang.set(qn("w:val"), language)
    default_lang.set(qn("w:eastAsia"), language)
    default_lang.set(qn("w:bidi"), language)

    for style in doc.styles:
        style_rpr = style._element.get_or_add_rPr()
        style_lang = style_rpr.find(qn("w:lang"))
        if style_lang is None:
            style_lang = OxmlElement("w:lang")
            style_rpr.append(style_lang)
        style_lang.set(qn("w:val"), language)
        style_lang.set(qn("w:eastAsia"), language)
        style_lang.set(qn("w:bidi"), language)

    settings = doc.settings.element
    theme_lang = settings.find(qn("w:themeFontLang"))
    if theme_lang is None:
        theme_lang = OxmlElement("w:themeFontLang")
        settings.append(theme_lang)
    theme_lang.set(qn("w:val"), language)
    theme_lang.set(qn("w:eastAsia"), language)
    theme_lang.set(qn("w:bidi"), language)


def configure_document(doc: Document) -> None:
    set_document_language(doc)
    section = doc.sections[0]
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.68)
    section.bottom_margin = Inches(0.67)
    section.left_margin = Inches(0.72)
    section.right_margin = Inches(0.72)
    section.header_distance = Inches(0.28)
    section.footer_distance = Inches(0.3)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Liberation Sans"
    normal.font.size = Pt(10.3)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Liberation Sans")
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE

    body = styles["Body Text"]
    body.font.name = "Liberation Sans"
    body.font.size = Pt(10.3)
    body.font.color.rgb = RGBColor.from_string(INK)
    body._element.rPr.rFonts.set(qn("w:eastAsia"), "Liberation Sans")
    body.paragraph_format.space_after = Pt(5)
    body.paragraph_format.line_spacing = 1.08
    body.paragraph_format.left_indent = Inches(0)
    body.paragraph_format.right_indent = Inches(0)
    body.paragraph_format.first_line_indent = Inches(0)

    for level, size, before, after in ((1, 17, 8, 5), (2, 13, 7, 4), (3, 11, 5, 3)):
        style = styles[f"Heading {level}"]
        style.font.name = "Liberation Sans"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(INK)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Liberation Sans")
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    caption = styles["Caption"]
    caption.font.name = "Liberation Sans"
    caption.font.size = Pt(8.2)
    caption.font.italic = False
    caption.font.color.rgb = RGBColor.from_string(MUTED)
    caption._element.rPr.rFonts.set(qn("w:eastAsia"), "Liberation Sans")
    caption.paragraph_format.space_after = Pt(5)

    for style_name in ("List Bullet", "List Bullet 2"):
        style = styles[style_name]
        style.font.name = "Liberation Sans"
        style.font.size = Pt(10.1)
        style.paragraph_format.space_after = Pt(3)

    header = section.header.paragraphs[0]
    header.text = "Informe final | VAE no equivariante vs. VAE SO(2)"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for run in header.runs:
        run.font.name = "Liberation Sans"
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor.from_string(MUTED)
    add_page_number(section.footer.paragraphs[0])
    for run in section.footer.paragraphs[0].runs:
        run.font.name = "Liberation Sans"
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor.from_string(MUTED)

    props = doc.core_properties
    props.title = "Informe final de resultados experimentales"
    props.subject = es(
        "Comparacion entre VAE no equivariante y VAE equivariante continuo SO(2)",
    )
    props.author = "Maximiliano Garavito Chtefan"
    props.language = "es-CO"
    props.keywords = es("VAE; SO(2); histopatologia; reconstruccion; MIL; tejido")
    props.comments = "Generado desde artefactos experimentales aceptados."


def add_cover(doc: Document) -> None:
    for _ in range(3):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("INFORME FINAL DE\nRESULTADOS EXPERIMENTALES")
    run.bold = True
    run.font.name = "Liberation Sans"
    run.font.size = Pt(26)
    run.font.color.rgb = RGBColor.from_string(INK)
    p.paragraph_format.space_after = Pt(24)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(
        es(
            "Comparacion entre un VAE no equivariante y un VAE\nequivariante continuo SO(2) para histopatologia",
        ),
    )
    run.font.name = "Liberation Sans"
    run.font.size = Pt(16)
    run.font.color.rgb = RGBColor.from_string(MUTED)
    p.paragraph_format.space_after = Pt(52)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("Proyecto de tesis\n").bold = True
    p.add_run("Maximiliano Garavito Chtefan\n")
    p.add_run("Universidad Industrial de Santander")
    for run in p.runs:
        run.font.name = "Liberation Sans"
        run.font.size = Pt(12)
        run.font.color.rgb = RGBColor.from_string(INK)
    p.paragraph_format.space_after = Pt(70)
    p = doc.add_paragraph("9 de septiembre de 2026")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in p.runs:
        run.font.size = Pt(10)
        run.font.color.rgb = RGBColor.from_string(MUTED)
    doc.add_page_break()


def build_report(repo: Path, output_docx: Path, work_dir: Path) -> None:
    fixed = read_json(
        repo
        / "runs/local/professor_metrics_v1/metrics/reconstruction_fixed25_summary.json",
    )
    recon = read_json(
        repo
        / "runs/local/vae_test_reconstruction_scored_v1/metrics/overall_summary.json",
    )
    mil_normal = read_json(
        repo
        / "runs/local/ubc_ocean_mil_test_scored_v1/normal_vae/scored_predictions_and_metrics.json",
    )
    mil_so2 = read_json(
        repo
        / "runs/local/ubc_ocean_mil_test_scored_v1/so2_vae/scored_predictions_and_metrics.json",
    )
    mil_pair = read_json(
        repo / "runs/local/ubc_ocean_mil_test_scored_v1/paired_bootstrap.json",
    )
    tissue = read_json(
        repo / "runs/local/tissue_test_scored_v1/label_efficiency_table.json",
    )
    mil_training_root = (
        repo / "runs/kaggle/ubc_ocean_mil_training_v7/spec0036_mil_training"
    )
    mil_train = {
        branch: read_csv(mil_training_root / branch / "metrics/train.csv")
        for branch in ("normal_vae", "so2_vae")
    }
    mil_validation = {
        branch: read_csv(mil_training_root / branch / "metrics/validation.csv")
        for branch in ("normal_vae", "so2_vae")
    }
    mil_summaries = {
        branch: read_json(mil_training_root / branch / "final_summary.json")
        for branch in ("normal_vae", "so2_vae")
    }
    tissue_training_root = (
        repo / "runs/kaggle/tissue_label_efficiency_training_v0003/"
        "tissue_label_efficiency_training/branches"
    )
    tissue_validation = []
    for budget in (250, 500, 1000, 2500, 5671):
        row: dict[str, float] = {"budget": float(budget)}
        for branch in ("normal_vae", "so2_vae"):
            result = read_json(
                tissue_training_root
                / branch
                / f"budget_{budget:04d}_per_class/result.json",
            )
            row[branch] = result["branches"][branch]["best"]["macro_f1"]
        tissue_validation.append(row)

    work_dir.mkdir(parents=True, exist_ok=True)
    output_docx.parent.mkdir(parents=True, exist_ok=True)
    mil_chart = work_dir / "mil_test_summary.png"
    mil_class_chart = work_dir / "mil_test_class_detail.png"
    tissue_chart = work_dir / "tissue_label_efficiency.png"
    tissue_class_chart = work_dir / "tissue_label_efficiency_by_class.png"
    supervised_development_chart = work_dir / "supervised_development.png"
    draw_mil_chart(mil_normal, mil_so2, mil_chart)
    draw_wsi_class_chart(mil_normal, mil_so2, mil_class_chart)
    draw_tissue_chart(tissue, tissue_chart)
    draw_tissue_class_chart(tissue, tissue_class_chart)
    draw_supervised_development_chart(
        mil_train,
        mil_validation,
        mil_summaries,
        tissue_validation,
        supervised_development_chart,
    )

    professor_figures = repo / "runs/local/professor_metrics_v1/figures"
    recon_figures = repo / "runs/local/vae_test_reconstruction_scored_v1/figures"
    legacy_orbit_figures = repo / "runs/local/frozen_vae_rotation_orbits"
    rotation_root = (
        repo
        / "runs/local/corrected_rotation_geometry_v1/corrected_rotation_geometry_v1"
    )
    rotation_figures = rotation_root / "figures"
    rotation_summary_path = rotation_root / "rotation_geometry_summary.json"
    rotation_addendum_path = (
        repo / "docs/data/spec0050_rotation_geometry_review_addendum.json"
    )
    decoded_root = (
        repo / "runs/local/decoded_latent_transform_v1/decoded_latent_transform_v1"
    )
    decoded_summary_path = decoded_root / "decoded_latent_transform_summary.json"
    decoded_manifest_path = decoded_root / "manifest.json"
    decoded_addendum_path = (
        repo / "docs/data/spec0051_decoded_transform_review_addendum.json"
    )
    required = [
        professor_figures / "training_dashboard.png",
        professor_figures / "metrics_boxplots_fixed25.png",
        professor_figures / "reconstructions_fixed25.png",
        professor_figures / "rotated_input_vs_latent.png",
        rotation_figures / "01-corrected-all25-orbits.png",
        rotation_figures / "02-population-input-controls.png",
        rotation_figures / "03-local-pca-ranks-00-12.png",
        rotation_figures / "03b-local-pca-pc1-pc6.png",
        rotation_figures / "04-pc1-pc6-harmonics.png",
        rotation_figures / "05-all48-f1-summary.png",
        rotation_figures / "06-shared-generator-vs-dimension.png",
        rotation_figures / "07-canonicalization-factorization.png",
        rotation_summary_path,
        rotation_root / "selected_arrays.npz",
        rotation_root / "manifest.json",
        rotation_addendum_path,
        decoded_summary_path,
        decoded_manifest_path,
        decoded_addendum_path,
        legacy_orbit_figures / "05-paper-style-latent-pca.png",
        recon_figures / "patch_metric_boxplots.png",
        recon_figures / "paired_wsi_mae.png",
    ]
    missing = [str(path) for path in required if not path.exists()]
    if missing:
        raise FileNotFoundError("Faltan figuras aceptadas:\n" + "\n".join(missing))
    if sha256_file(rotation_summary_path) != ROTATION_GEOMETRY_SUMMARY_SHA256:
        message = "El resumen de geometría rotacional cambió después de la revisión"
        raise ValueError(message)
    if sha256_file(rotation_addendum_path) != ROTATION_GEOMETRY_ADDENDUM_SHA256:
        message = "El addendum de revisión geométrica cambió después de la revisión"
        raise ValueError(message)
    rotation_geometry = validate_rotation_geometry(
        read_json(rotation_summary_path),
        read_json(rotation_addendum_path),
    )
    for relative_path, expected_hash in rotation_geometry["output_hashes"].items():
        artifact_path = rotation_root / relative_path
        if sha256_file(artifact_path) != expected_hash:
            message = f"Hash geométrico inesperado: {relative_path}"
            raise ValueError(message)
    if sha256_file(decoded_summary_path) != DECODED_TRANSFORM_SUMMARY_SHA256:
        message = "El resumen de Spec 0051 cambió después de la revisión"
        raise ValueError(message)
    if sha256_file(decoded_manifest_path) != DECODED_TRANSFORM_MANIFEST_SHA256:
        message = "El manifiesto de Spec 0051 cambió después de la revisión"
        raise ValueError(message)
    if sha256_file(decoded_addendum_path) != DECODED_TRANSFORM_ADDENDUM_SHA256:
        message = "El addendum de Spec 0051 cambió después de la revisión"
        raise ValueError(message)
    decoded_addendum = read_json(decoded_addendum_path)
    decoded_transform = validate_decoded_transform(
        read_json(decoded_summary_path),
        decoded_addendum,
    )
    for relative_path, expected_hash in decoded_transform["output_hashes"].items():
        artifact_path = decoded_root / relative_path
        if sha256_file(artifact_path) != expected_hash:
            message = f"Hash de Spec 0051 inesperado: {relative_path}"
            raise ValueError(message)
    corrected_llr = rotation_geometry["comparisons"]["local_linearity_ratio"]
    corrected_cv = rotation_geometry["comparisons"]["step_size_cv"]
    corrected_subsampling = rotation_geometry["comparisons"]["subsampling"]
    decoded_dense_review = decoded_addendum["dense_rotation"]
    decoded_action_review = decoded_dense_review["action_ratio_median"]
    decoded_canonical_review = decoded_dense_review["canonical_ratio_median"]
    decoded_input_review = decoded_dense_review[
        "end_to_end_input_commutation_ratio_median"
    ]
    decoded_quarter_review = decoded_addendum["exact_quarter_rotations"]
    decoded_action_interval = decoded_action_review["descriptive_interval"]
    decoded_canonical_interval = decoded_canonical_review["descriptive_interval"]
    decoded_action_quarter = decoded_quarter_review["aggregate_action_ratio_median"]
    decoded_canonical_quarter = decoded_quarter_review[
        "aggregate_canonical_ratio_median"
    ]
    decoded_robustness = decoded_dense_review["so2_robustness"]
    corrected_harmonics = rotation_geometry["comparisons"]["harmonics"]
    rotation_decisions = rotation_geometry["decisions"]
    orbit_population_sheet_a = work_dir / "orbit_population_patches_00_14.png"
    orbit_population_sheet_b = work_dir / "orbit_population_patches_15_24.png"
    draw_orbit_population_report_sheets(
        rotation_figures / "01-corrected-all25-orbits.png",
        orbit_population_sheet_a,
        orbit_population_sheet_b,
    )
    population_controls_figure = work_dir / "population_input_controls_fixed25.png"
    draw_population_controls_report_figure(
        rotation_figures / "02-population-input-controls.png",
        population_controls_figure,
    )
    decoded_population_figure = work_dir / "decoded_transform_population.png"
    decoded_exact_figure = work_dir / "decoded_transform_exact.png"
    draw_decoded_transform_population_figure(
        decoded_transform,
        decoded_population_figure,
    )
    draw_decoded_transform_exact_figure(decoded_transform, decoded_exact_figure)

    doc = Document()
    configure_document(doc)
    add_cover(doc)

    add_heading(doc, "1. Resumen ejecutivo", 1)
    add_body(
        doc,
        "Este informe responde las solicitudes de evaluacion registradas por el profesor en los issues #2, #3, #4 y #6: curvas y graficas de resultados, MAE, MSE, PSNR, SSIM, desviacion estandar, tamano muestral, diagramas de caja, 25 reconstrucciones fijas, rotaciones de 90/180/270 grados y una visualizacion PCA del espacio latente para ambos modelos.",
    )
    add_body(
        doc,
        "Conclusion principal. La evidencia completa es mixta y no permite declarar un ganador universal.",
        bold_prefix="Conclusion principal.",
    )
    add_bullet(
        doc,
        "Reconstruccion sellada: el VAE no equivariante obtuvo un MAE promedio ligeramente menor. La diferencia primaria normal menos SO(2) fue -0,001358, pero su IC 95% por bootstrap de WSI [-0,002213; 0,000046] incluye cero; no se establecio una diferencia.",
    )
    add_bullet(
        doc,
        "Diagnostico WSI: SO(2) fue numericamente mejor en F1 macro (0,483 frente a 0,389) y exactitud (0,609 frente a 0,522), pero el IC 95% de la diferencia de F1 macro tambien incluye cero.",
    )
    add_bullet(
        doc,
        "Tejido: SO(2) tuvo un F1 macro numericamente mayor en los cuatro presupuestos bajos; solo 500 etiquetas por clase tuvo un intervalo simultaneo que excluye cero. El area global bajo la curva de aprendizaje (AULC) fue inconclusa y, en el presupuesto maximo, el modelo normal tuvo mayor F1 macro.",
    )
    add_bullet(
        doc,
        "Geometria rotacional post-hoc: se reprodujo y corrigio un cambio de signo que unia orientaciones opuestas en 90 y 270 grados. El aparente patron previo queda retirado. A 1 grado, SO(2) tuvo menor razon de linealidad local normalizada por la entrada en 25/25 parches (medianas 0,601 frente a 0,649), pero el efecto fue 7,38% y no sobrevivio la robustez preespecificada: a 5 grados solo ocurrio en 8/25. No se demostro una accion reducida compartida ni una factorizacion local contenido--pose.",
    )
    add_bullet(
        doc,
        f"Acción en el decodificador: al transformar espacialmente el embedding final antes de decodificar, SO(2) redujo el error frente al VAE normal en {decoded_action_review['so2_favorable_patches']}/{FIXED25_COUNT} parches para el barrido de 5 grados. En rotaciones exactas de 90, 180 y 270 grados, su razón de acción fue {es_decimal(float(decoded_action_quarter['so2']), 6)} frente a {es_decimal(float(decoded_action_quarter['normal']), 3)} normal. Esto verifica empíricamente una acción C4 compartida y sin parámetros ajustados en el decodificador de este checkpoint; no prueba equivarianza continua del encoder ni factorización contenido--pose.",
    )

    add_heading(doc, "Lectura para la presentacion", 2)
    add_table(
        doc,
        ["Pregunta", "Respuesta respaldada"],
        [
            [
                "¿Reconstruye mejor SO(2)?",
                "No. El modelo normal es ligeramente mejor en promedio, sin diferencia primaria establecida.",
            ],
            [
                "¿Ayuda en tareas posteriores?",
                "Hay senales favorables a SO(2) en WSI y en bajo etiquetado, pero no una superioridad general.",
            ],
            [
                "¿La equivarianza es visible?",
                "Sí en el decodificador para rotaciones exactas C4: transformar el embedding rota la reconstrucción de SO(2) casi exactamente. No se verificó una acción continua compartida en mu ni campos F1 limpios.",
            ],
            [
                "¿Falta experimentar?",
                "Los issues quedan respondidos. La geometria nueva es validacion fija exploratoria y no justifica afirmar un producto global o disentanglement.",
            ],
        ],
        widths=[2.1, 4.75],
        font_size=9,
    )

    add_heading(doc, "2. Pregunta experimental y comparacion", 1)
    add_body(
        doc,
        "Se compararon dos codificadores variacionales entrenados sobre el mismo dominio histopatologico: una linea base convolucional no equivariante y una arquitectura con equivarianza continua SO(2). Ambos completaron 60.000 actualizaciones. Las particiones, los checkpoints y el cierre del entrenamiento precedieron al test; sin embargo, la agregacion exacta y el bootstrap de reconstruccion se formalizaron despues de abrir las etiquetas de los tests posteriores. Esos resultados no se usaron para cambiar modelos ni seleccionar checkpoints.",
    )
    add_table(
        doc,
        ["Propiedad", NORMAL_LABEL, SO2_LABEL],
        [
            [
                "Simetria arquitectonica",
                "Sin restriccion rotacional explicita",
                "Equivarianza continua SO(2)",
            ],
            ["Actualizaciones", "60.000", "60.000"],
            ["Parametros aprendibles", "3.958.435", "1.180.035"],
            ["Tamano relativo", "100%", "29,8% del modelo normal"],
        ],
        widths=[1.7, 2.55, 2.55],
    )
    add_body(
        doc,
        "La reduccion aproximada de 70,2% en parametros es relevante para eficiencia, pero tambien es una limitacion causal: la comparacion cambia simultaneamente la simetria y la capacidad del modelo.",
    )

    add_heading(doc, "Implementacion SO(2) verificada antes del experimento", 2)
    add_body(
        doc,
        "La implementacion continua SO(2) fue validada por componentes y como VAE completo antes del entrenamiento sobre datos reales. Esta verificacion responde la parte de implementacion del issue #6.",
    )
    add_table(
        doc,
        ["Componente", "Evidencia de verificacion"],
        [
            [
                "No linealidades y normalizacion",
                "34 compuertas radiales F0/F1 y 40 normalizaciones: aritmetica FP32, gradientes y actualizaciones finitas.",
            ],
            [
                "Remuestreo y upsampling",
                "Seis downsamplers y seis upsamplers fijos; formas 256-128-64-32 y 32-64-128-256 verificadas.",
            ],
            [
                "Muestreo VAE y estadisticas latentes",
                "mu, logvar, z con epsilon controlado, decodificador y reconstruccion comprobados en angulos cardinales y no cardinales; mu tiene forma (B,16,32,32).",
            ],
            [
                "Referencia y runtime",
                "escnn se uso solo como oraculo de pruebas focalizadas; no es dependencia del runtime. La implementacion propia compilable pasó el prelaunch batch-25 dual T4 sin skips AMP ni recompilaciones.",
            ],
        ],
        widths=[2.05, 4.65],
        font_size=8.4,
    )

    add_heading(doc, "3. Protocolo de evaluacion", 1)
    add_table(
        doc,
        ["Capa de evidencia", "Poblacion", "Unidad", "Uso"],
        [
            [
                "Validacion fija",
                "25 parches predeterminados",
                "Parche",
                "Descripcion y control visual; no es test sellado",
            ],
            [
                "Reconstruccion sellada",
                "67.138 parches de 23 WSI",
                "Parche; bootstrap por WSI",
                "Comparacion primaria de MAE y endpoints secundarios",
            ],
            [
                "Diagnostico WSI sellado",
                "23 WSI",
                "WSI",
                "F1 macro, exactitudes y entropia cruzada",
            ],
            [
                "Tejido sellado",
                "31.572 parches de 23 WSI",
                "Parche; remuestreo por WSI",
                "Eficiencia de etiquetas en cinco presupuestos",
            ],
        ],
        widths=[1.35, 1.7, 1.65, 2.1],
        font_size=8.4,
    )
    add_body(
        doc,
        "Regla de inferencia. En reconstruccion, MAE normalizado fue el endpoint confirmatorio. MSE, PSNR y SSIM fueron secundarios. Las desviaciones estandar de parches describen dispersion y no sustituyen los intervalos de incertidumbre por WSI.",
        bold_prefix="Regla de inferencia.",
    )
    add_body(
        doc,
        "Direccion de contrastes. Salvo que se indique lo contrario, las diferencias se informan como normal menos SO(2): valores negativos favorecen al modelo normal en errores (MAE/MSE) y favorecen a SO(2) en metricas donde un valor mayor es mejor (F1/exactitud).",
        bold_prefix="Direccion de contrastes.",
    )

    doc.add_page_break()
    add_heading(doc, "4. Entrenamiento y validacion fija de 25 parches", 1)
    add_body(
        doc,
        "La validacion fija permite comprobar convergencia, estabilidad y calidad visual con exactamente los mismos parches para ambos modelos. Sus resultados son descriptivos; no reemplazan la evaluacion sellada.",
    )
    add_figure(
        doc,
        professor_figures / "training_dashboard.png",
        6.75,
        "Figura 1. Panel de entrenamiento hasta 60.000 actualizaciones: objetivo, Charbonnier/L1, 1-SSIM, PSNR, diagnostico latente de equivarianza y tasa de aprendizaje. Validacion fija de 25 parches; no es test sellado.",
        "Panel de seis graficas de entrenamiento para los modelos normal y SO2 hasta 60000 actualizaciones.",
    )

    normal_fixed = fixed["models"]["normal"]
    so2_fixed = fixed["models"]["so2"]
    metric_rows = []
    for label, key in (
        ("MAE (menor es mejor)", "mae_norm"),
        ("MSE (menor es mejor)", "mse_norm"),
        ("PSNR dB (mayor es mejor)", "psnr_img_db"),
        ("SSIM (mayor es mejor)", "ssim_img"),
    ):
        nstat, sstat = normal_fixed[key], so2_fixed[key]
        decimals = 3 if key == "psnr_img_db" else 4
        metric_rows.append([
            label,
            f"{nstat['mean']:.{decimals}f} ({nstat['std_population']:.{decimals}f})",
            f"{sstat['mean']:.{decimals}f} ({sstat['std_population']:.{decimals}f})",
            "25 / 25",
        ])
    add_table(
        doc,
        ["Metrica", "Normal: media (DEp)", "SO(2): media (DEp)", "n"],
        metric_rows,
        widths=[2.2, 1.8, 1.8, 0.9],
        font_size=8.7,
    )
    add_body(
        doc,
        "En esta muestra descriptiva, el modelo normal tuvo menor MAE/MSE y mayor PSNR/SSIM. No se aplico inferencia estadistica a la validacion fija.",
    )
    add_figure(
        doc,
        professor_figures / "metrics_boxplots_fixed25.png",
        6.65,
        "Figura 2. Diagramas de caja solicitados para MAE, MSE, PSNR y SSIM en los 25 parches fijos. Cada caja resume n=25 por modelo; DEp significa desviacion estandar poblacional.",
        "Cuatro diagramas de caja comparan MAE, MSE, PSNR y SSIM en 25 parches para los dos modelos.",
    )

    doc.add_page_break()
    add_heading(doc, "5. Reconstrucciones cualitativas fijas", 1)
    add_body(
        doc,
        "Las reconstrucciones permiten verificar que ambos modelos preservan estructura y color generales. La inspeccion visual complementa, pero no reemplaza, las metricas cuantitativas.",
    )
    add_figure(
        doc,
        professor_figures / "reconstructions_fixed25.png",
        6.9,
        "Figura 3. Originales y reconstrucciones de los 25 parches predeterminados para ambos modelos. Poblacion: validacion fija; n=25; seleccion fijada antes del analisis.",
        "Mosaico de 25 parches originales junto con sus reconstrucciones por el VAE normal y el VAE SO2.",
    )
    add_heading(doc, "Cómo leer esta figura", 2)
    add_bullet(
        doc,
        "Buscar perdida de bordes celulares, suavizado excesivo, cambios de color y artefactos repetitivos.",
    )
    add_bullet(
        doc,
        "No elegir un modelo por un parche aislado: las metricas selladas agregan 67.138 parches.",
    )
    add_bullet(
        doc,
        "La seleccion fija evita escoger ejemplos favorables despues de observar los resultados.",
    )

    doc.add_page_break()
    add_heading(doc, "6. Reconstruccion en el test sellado", 1)
    add_body(
        doc,
        "La evaluacion principal amplio la comparacion a 67.138 parches procedentes de 23 WSI. Los intervalos se calcularon con 10.000 remuestras por conglomerado WSI; cubren variacion entre WSI, no nuevas semillas de entrenamiento ni una nueva seleccion de parches.",
    )
    recon_rows = []
    keys = (
        ("MAE", "mae_norm"),
        ("MSE", "mse_norm"),
        ("PSNR (dB)", "psnr_img"),
        ("SSIM", "ssim_img"),
    )
    for label, key in keys:
        nstat = recon["branches"]["normal"]["patch_distribution"][key]
        sstat = recon["branches"]["so2"]["patch_distribution"][key]
        decimals = 3 if key == "psnr_img" else 4
        delta = recon["paired_normal_minus_so2"][key]["point_difference"]
        recon_rows.append([
            label,
            f"{nstat['mean']:.{decimals}f} ({nstat['population_sd']:.{decimals}f})",
            f"{sstat['mean']:.{decimals}f} ({sstat['population_sd']:.{decimals}f})",
            f"{delta:+.{decimals}f}",
        ])
    add_table(
        doc,
        ["Metrica", "Normal: media (DEp)", "SO(2): media (DEp)", "Normal - SO(2)"],
        recon_rows,
        widths=[1.55, 1.9, 1.9, 1.35],
        font_size=8.6,
    )
    add_body(
        doc,
        "Resultado primario. MAE: 0,062872 para el modelo normal y 0,064230 para SO(2). Diferencia normal menos SO(2) = -0,001358; IC 95% [-0,002213; 0,000046]. Como el intervalo incluye cero, no se establecio una diferencia de reconstruccion primaria.",
        bold_prefix="Resultado primario.",
    )
    add_body(
        doc,
        "Endpoints secundarios. PSNR y SSIM favorecen descriptivamente al modelo normal, y sus intervalos exploratorios no cruzan cero. Se mantienen como resultados secundarios: no sustituyen la conclusion del endpoint primario.",
        bold_prefix="Endpoints secundarios.",
    )
    add_figure(
        doc,
        recon_figures / "patch_metric_boxplots.png",
        6.6,
        "Figura 4. Distribuciones por parche en el test sellado: 67.138 parches por modelo, agrupados en 23 WSI. Las cajas describen dispersion de parches; no son intervalos inferenciales. La tabla superior proporciona la escala numerica.",
        "Diagramas de caja del test sellado comparan las cuatro metricas de reconstruccion entre ambos modelos.",
    )

    doc.add_page_break()
    add_heading(doc, "7. Consistencia entre WSI", 1)
    add_body(
        doc,
        "La vista pareada por WSI evita que el gran numero de parches oculte la unidad biologica. En 22 de los 23 WSI, el MAE medio del modelo normal fue menor; un WSI mostro el patron contrario. Aun asi, la incertidumbre del estimando primario ponderado por parches alcanza el cero.",
    )
    add_figure(
        doc,
        recon_figures / "paired_wsi_mae.png",
        6.8,
        "Figura 5. MAE medio pareado por WSI en el test sellado (n=23 WSI). Cada linea conecta exactamente el mismo WSI bajo ambos modelos.",
        "Grafica pareada con 23 lineas que conectan el MAE de cada WSI para el modelo normal y el modelo SO2.",
    )
    add_heading(doc, "Interpretacion", 2)
    add_body(
        doc,
        "El patron por WSI sugiere una pequena ventaja de fidelidad para el modelo normal, coherente con los endpoints secundarios. Sin embargo, el objetivo del experimento no era optimizar solo pixeles: tambien se evaluo si las representaciones resultaban utiles para diagnostico y tejido.",
    )

    doc.add_page_break()
    add_heading(doc, "8. Geometría de la acción rotacional en el espacio latente", 1)
    add_body(
        doc,
        "Este análisis es un trabajo post-hoc, exploratorio y exclusivo de los 25 parches fijos de validación. No accedió al test sellado, no reentrenó los VAE y no cambió checkpoints. Su pregunta principal no fue si cada parche dibuja un ciclo —resultado casi trivial para un codificador continuo— sino si una misma acción rotacional reducida transfiere entre parches y permite separar aproximadamente contenido y pose.",
    )
    add_heading(doc, "Control exacto a cuartos de vuelta", 2)
    add_body(
        doc,
        "La comparación solicitada originalmente a 90, 180 y 270 grados usa exclusivamente torch.rot90 y permanece válida. La figura conserva el parche predeterminado de rango 12 y compara reconstruir la entrada rotada con transformar espacialmente el latente antes de decodificar.",
    )
    add_figure(
        doc,
        professor_figures / "rotated_input_vs_latent.png",
        6.5,
        "Figura 6. Comparacion cualitativa a 0, 90, 180 y 270 grados: verdad objetivo, reconstruccion de la entrada rotada y reconstruccion despues de transformar el latente. n=1, selector rank 12 del fixed-25; no es resultado poblacional ni test sellado.",
        "Cuadricula de rotaciones que compara imagen objetivo, reconstruccion de entrada rotada y reconstruccion de latente rotado para ambos modelos.",
    )
    add_heading(doc, "Resultado cuantitativo complementario", 2)
    add_body(
        doc,
        "En el único parche mostrado, la diferencia MAE entre las dos rutas de decodificación fue 0,1002/0,1097/0,0988 para el VAE normal y 0,0118/0,0130/0,0109 para SO(2) a 90/180/270 grados. Esto respalda una mayor consistencia entre rutas en ese ejemplo, pero n=1 impide generalizarlo.",
    )
    add_body(
        doc,
        "En el proxy enmascarado de residuo latente para los 25 parches, menor es mejor. Las medianas aceptadas fueron 1,327492/1,303358/1,326518 para el modelo normal y 1,472552/1,417673/1,486783 para SO(2). Este control exacto no favorece a SO(2) en la salida posterior final.",
    )
    add_heading(doc, "Acción decodificada y comparación pareada", 2)
    add_body(
        doc,
        "El experimento complementario de Spec 0051 evaluó dos rutas sobre los mismos 25 parches. La razón de acción compara D(rho_T mu(x)) con T D(mu(x)); la razón de canonicalización compara D(rho_T^-1 mu(Tx)) con D(mu(x)). Cada denominador es el error de dejar el embedding sin la transformación correspondiente. Por tanto, 1,00 representa la línea base de no actuar y 0,50 exige conservar como máximo la mitad de ese error. Este 0,50 es un criterio funcional absoluto fijado antes de ver la salida; no decide qué VAE es mejor. La comparación entre modelos es pareada sobre los mismos parches y se resume también por sus 16 WSI de origen.",
    )
    add_body(
        doc,
        f"En el barrido 0, 5, ..., 355 grados, SO(2) tuvo menor error que el modelo normal en {decoded_action_review['so2_favorable_patches']}/{FIXED25_COUNT} parches y {decoded_action_review['so2_favorable_wsi_clusters']}/16 medianas por WSI. La razón de acción mediana fue {es_decimal(float(decoded_action_review['so2']), 3)} SO(2) frente a {es_decimal(float(decoded_action_review['normal']), 3)} normal, una reducción relativa de {es_decimal(100.0 * float(decoded_action_review['relative_so2_reduction']), 2)}%; la diferencia pareada SO(2) menos normal por bootstrap descriptivo de WSI fue {es_decimal(float(decoded_action_review['wsi_cluster_median_difference_so2_minus_normal']), 3)} [{es_decimal(float(decoded_action_interval[0]), 3)}; {es_decimal(float(decoded_action_interval[1]), 3)}]. Para canonicalización, las medianas fueron {es_decimal(float(decoded_canonical_review['so2']), 3)} frente a {es_decimal(float(decoded_canonical_review['normal']), 3)}, reducción de {es_decimal(100.0 * float(decoded_canonical_review['relative_so2_reduction']), 2)}%, con diferencia {es_decimal(float(decoded_canonical_review['wsi_cluster_median_difference_so2_minus_normal']), 3)} [{es_decimal(float(decoded_canonical_interval[0]), 3)}; {es_decimal(float(decoded_canonical_interval[1]), 3)}]. El conjunto constituye una ventaja comparativa consistente de SO(2) en la consistencia de rutas decodificadas. H1 y H2 completos permanecen no respaldados porque sus reglas conjuntivas exigían además una mediana SO(2) <=0,50; ambos valores quedaron ligeramente por encima. Ese criterio absoluto no anula la comparación pareada.",
    )
    add_figure(
        doc,
        decoded_population_figure,
        6.55,
        f"Figura 6a. Razones pareadas de acción espacial y canonicalización decodificadas para los 25 parches fijos de validación, con rotaciones 0, 5, ..., 355 grados. Cada línea une el mismo parche. SO(2) es menor en {decoded_action_review['so2_favorable_patches']}/{FIXED25_COUNT} pares en ambas métricas. El eje parte de cero y la línea gris discontinua marca el criterio funcional absoluto 0,50; los intervalos resumen primero por WSI y remuestrean las 16 WSI, sin inferencia poblacional.",
        "Dos gráficos pareados muestran menor razón de error para SO2 en los 25 parches; una línea horizontal identifica el criterio absoluto 0,50.",
    )
    add_body(
        doc,
        f"El resultado exacto es mucho más fuerte. Agregando 90, 180 y 270 grados, la razón de acción fue {es_decimal(float(decoded_action_quarter['so2']), 8)} SO(2) frente a {es_decimal(float(decoded_action_quarter['normal']), 6)} normal, con error absoluto numéricamente cercano a cero y señal de pose no nula. La canonicalización decodificada fue {es_decimal(float(decoded_canonical_quarter['so2']), 6)} frente a {es_decimal(float(decoded_canonical_quarter['normal']), 6)}. Así, el decodificador SO(2) realiza numéricamente la misma acción espacial C4 para todos los parches fijos, sin ajustar parámetros por parche. Esto reconcilia la observación visual con el residuo desfavorable de mu: el posterior crudo no coincide mejor bajo ese proxy, pero las dos rutas convergen después de decodificar. El experimento no identifica un espacio nulo del decodificador.",
    )
    add_body(
        doc,
        f"Los ángulos no cardinales siguen siendo aproximados: al excluir vecindades de 5 grados alrededor de los cardinales, las razones SO(2) suben a {es_decimal(float(decoded_robustness['exclude_cardinal_neighborhoods_action_canonical'][0]), 3)} para acción y {es_decimal(float(decoded_robustness['exclude_cardinal_neighborhoods_action_canonical'][1]), 3)} para canonicalización. Además, la ruta completa D(mu(Tx)) frente a T D(mu(x)) favoreció al VAE normal en {FIXED25_COUNT - int(decoded_input_review['so2_favorable_patches'])}/{FIXED25_COUNT} parches, con medianas {es_decimal(float(decoded_input_review['normal']), 3)} normal y {es_decimal(float(decoded_input_review['so2']), 3)} SO(2). Por ello, el hallazgo no equivale a una superioridad general del encoder o del autoencoder.",
    )
    add_body(
        doc,
        "En los cuatro flips exactos, SO(2) tuvo razones descriptivamente menores, pero sólo la reflexión respecto a la diagonal principal cumplió simultáneamente los criterios de entrada, acción y canonicalización. Las reflexiones horizontal, vertical y antidiagonal no pasaron. No se demostró equivarianza completa D4 ni O(2).",
    )
    add_figure(
        doc,
        decoded_exact_figure,
        6.55,
        "Figura 6b. Medianas de acción latente, canonicalización y conmutación de entrada a salida para las siete transformaciones no identidad del grupo D4 exacto. Las rotaciones R90/R180/R270 muestran la acción C4 casi exacta del decodificador SO(2). Entre los flips, sólo Fd cumple la regla completa para ese checkpoint. La línea gris discontinua marca el criterio absoluto 0,50; n=25 parches fijos de validación.",
        "Tres gráficos de barras comparan los modelos en rotaciones exactas y cuatro reflexiones, con una referencia horizontal en 0,50.",
    )
    add_body(
        doc,
        "En estas métricas, los objetivos son transformaciones de la reconstrucción producida por cada modelo, no imágenes originales independientes. El análisis evalúa consistencia entre rutas de transformación y decodificación; no reemplaza las métricas selladas de fidelidad de reconstrucción.",
    )
    add_heading(doc, "Defecto de convención, reparación y alcance", 2)
    add_body(
        doc,
        "El barrido denso anterior mezclaba dos convenciones: los múltiplos de 90 grados usaban la orientación positiva de torch.rot90, mientras que la matriz afín no cardinal convergía a la orientación opuesta. Al unir puntos consecutivos aparecían los segmentos rectos observados cerca de 90 y 270 grados. Esa evidencia, incluidas las Figuras 7, 8a y 8b anteriores y la afirmación universal 25/25, queda explícitamente supersedida.",
    )
    add_body(
        doc,
        "La reparación usa una sola interpolación bilineal continua en todo el ciclo 0--359 y reserva torch.rot90 como control independiente. Fixtures asimétricos de impulso, flecha y borde dieron error cardinal exactamente cero en CPU y GPU; el RMS máximo a ambos lados de un cardinal fue 0,00000520, el error máximo de orientación 0,0193 grados y los errores de inversión/composición quedaron por debajo de 0,024, atribuibles a interpolación. En el parche 12, los pasos RMS de entrada 87--88, 88--89, 89--90, 90--91 y 91--92 fueron 0,1011/0,1025/0,1118/0,1129/0,1035: desapareció el salto previo de aproximadamente 0,24.",
    )

    add_heading(doc, "Órbitas por parche y ausencia de acción reducida", 2)
    add_figure(
        doc,
        rotation_figures / "03-local-pca-ranks-00-12.png",
        6.5,
        "Figura 7. Geometría PCA local corregida de los rangos 0 y 12, elegidos antes de ver la salida. Cada parche y modelo usa coordenadas propias; signos, ejes y escalas no se comparan entre paneles. En los 25 parches fijos, las seis primeras PC capturan medianas de solo 11,09% normal y 10,86% SO(2), de modo que estas proyecciones son descriptivas.",
        "Trayectorias tridimensionales y varianza explicada de las órbitas corregidas para los rangos preseleccionados 0 y 12.",
    )
    add_body(
        doc,
        "La PCA de una órbita completa estima un subespacio lineal que contiene variación rotacional; no identifica por sí sola el vector tangente ni una acción común. Los tangentes de órbita necesitaron medianas de 120/114 dimensiones para explicar 90% de su variación (normal/SO(2)), y los subespacios locales entre parches fueron casi ortogonales. No apareció un plano tangente crudo compartido de baja dimensión.",
    )
    add_body(
        doc,
        "Los mosaicos siguientes muestran los 25 ciclos ya corregidos. No se seleccionaron ejemplos por apariencia: se fijaron los rangos 0 y 12. Cada panel ajusta una PCA local con ejes isótropos; las métricas se calcularon sobre mu crudo dentro del disco latente de radio 14, no sobre la proyección.",
    )
    add_figure(
        doc,
        orbit_population_sheet_a,
        6.75,
        "Figura 8a. Órbitas corregidas de los parches fijos 00--14 a intervalos de 1 grado. Los segmentos largos espurios de la versión anterior ya no aparecen. Cada par usa una PCA propia y escala isótropa.",
        "Primera cuadrícula de órbitas PCA corregidas para los parches fijos 00 a 14.",
    )
    doc.add_page_break()
    add_figure(
        doc,
        orbit_population_sheet_b,
        6.75,
        "Figura 8b. Órbitas corregidas de los parches fijos 15--24 a intervalos de 1 grado. La forma puede compararse dentro de cada panel, no como un sistema de coordenadas común entre parches o modelos.",
        "Segunda cuadrícula de órbitas PCA corregidas para los parches fijos 15 a 24.",
    )

    add_heading(doc, "Regularidad y controles de la trayectoria de entrada", 2)
    add_body(
        doc,
        f"A 1 grado, la razón de linealidad local normalizada por la entrada fue menor para SO(2) en {corrected_llr['so2_lower_count']}/25: medianas {corrected_llr['so2_median']:.3f} frente a {corrected_llr['normal_median']:.3f} normal, una reducción relativa de 7,38%. El intervalo bootstrap descriptivo de la mediana pareada SO(2) menos normal fue [{corrected_llr['paired_bootstrap_so2_minus_normal']['percentile_95'][0]:.3f}; {corrected_llr['paired_bootstrap_so2_minus_normal']['percentile_95'][1]:.3f}]. Sin embargo, el umbral preregistrado exigía al menos 10% y robustez de escala: a 5 grados SO(2) fue menor sólo en {corrected_subsampling['5']['local_linearity_ratio']['so2_lower_count']}/25 y las medianas se invirtieron ligeramente ({corrected_subsampling['5']['local_linearity_ratio']['so2_median']:.3f} frente a {corrected_subsampling['5']['local_linearity_ratio']['normal_median']:.3f}). H1 no pasó.",
    )
    add_body(
        doc,
        f"El CV crudo del paso a 1 grado favoreció al modelo normal: medianas {corrected_cv['normal_raw_median']:.4f} normal y {corrected_cv['so2_raw_median']:.4f} SO(2), con SO(2) menor en sólo {corrected_cv['so2_lower_raw_count']}/25. Tras normalizar por el recorrido de entrada, la comparación cambió a {corrected_cv['normal_median']:.4f}/{corrected_cv['so2_median']:.4f} y 18/25. La dependencia de resolución, normalización e interpolación impide afirmar una ventaja general de regularidad.",
    )
    add_figure(
        doc,
        population_controls_figure,
        6.55,
        "Figura 8c. Métricas de los 25 parches fijos y controles de la trayectoria de entrada para 1, 2 y 5 grados, con y sin vecindades cardinales. La mejora descriptiva a 1 grado no superó el criterio preregistrado ni fue robusta a 5 grados.",
        "Comparación descriptiva de los 25 parches fijos: linealidad local, curvatura y uniformidad del paso con controles de la ruta de entrada.",
    )

    add_heading(doc, "Organización armónica y campos internos F1", 2)
    add_body(
        doc,
        f"La fracción mediana de la energía orbital cruda total capturada por m=1,...,6 dentro de PC1--PC6 fue {median(corrected_harmonics['normal_low_fraction']):.3f} normal y {median(corrected_harmonics['so2_low_fraction']):.3f} SO(2); SO(2) fue mayor en sólo {corrected_harmonics['so2_higher_count']}/25. El número efectivo de frecuencias fue {median(corrected_harmonics['normal_effective_count']):.3f}/{median(corrected_harmonics['so2_effective_count']):.3f}; la reducción de SO(2), aunque presente en 18/25, fue 3,7%, por debajo del 10% fijado. H2 no pasó y las seis PC sólo capturan cerca de 11% de la variación total.",
    )
    add_figure(
        doc,
        rotation_figures / "04-pc1-pc6-harmonics.png",
        6.45,
        "Figura 8d. Energía armónica de PC1--PC6 para los ejemplos preseleccionados. Las escalas de color son internas a cada panel; múltiples frecuencias pueden surgir legítimamente al rotar textura espacial. No apareció una concentración SO(2) específica robusta.",
        "Mapas de calor de potencia armónica entre frecuencias cero y seis para las seis primeras coordenadas PCA.",
    )
    add_body(
        doc,
        "Antes de mu_head, la arquitectura contiene 48 copias F0 y 48 copias F1; la salida posterior mu contiene 16 campos espaciales F0 y ninguna copia F1 explícita. Se midieron las 48 F1 mediante vectores complejos promediados y, para evitar cancelación espacial, mediante el campo completo compensando rotación interna y de coordenadas. Ninguna de las 48 cumplió simultáneamente pureza m=1, pendiente de fase, amplitud y residuo; el residuo completo mediano por parche fue 1,906 y el RMSE de fase mediano 1,671 radianes. H3 no pasó. Esto no invalida la construcción del tipo de campo, pero impide afirmar que el checkpoint entrenado verificó esa ley empíricamente.",
    )
    add_figure(
        doc,
        rotation_figures / "05-all48-f1-summary.png",
        6.45,
        "Figura 8e. Diagnóstico de las 48 copias internas F1. Resultado: 0/48 copias limpias bajo los umbrales fijados; el residuo de campo completo confirma que la fase inestable no se explica únicamente por cancelación del promedio espacial.",
        "Resumen por copia F1 de pureza armónica, error de fase, estabilidad de amplitud y residuo de equivarianza del campo completo.",
    )

    add_heading(doc, "Acción compartida y factorización", 2)
    add_body(
        doc,
        "La partición derivada de hashes fijó 17 parches para ajuste y 8 para evaluación; se ajustó con ángulos múltiplos de 5 grados y se evaluaron los demás. Entre d=1,...,6, la regla fijada seleccionó d=1, que es necesariamente una acción ortogonal continua nula. En held-out, el NRMSE/R2 fue 1,420/-1,017 normal y 1,450/-1,103 SO(2); ningún parche (0/8 en ambos modelos) mejoró frente a identidad. Incluso d=6 capturó sólo 0,15%/0,24% de la varianza held-out. Por tanto, H4 no pasó y no se demostró una acción reducida compartida. La ley de grupo casi exacta de exp(theta A) es una propiedad numérica impuesta por la parametrización, no evidencia de que A prediga las órbitas.",
    )
    add_figure(
        doc,
        rotation_figures / "06-shared-generator-vs-dimension.png",
        6.55,
        "Figura 8f. Rendimiento held-out del generador compartido frente a dimensión. La dimensión seleccionada fue d=1, línea base nula; ninguna dimensión mostró una ventaja específica de SO(2) sobre identidad o el control de ángulos barajados.",
        "Curvas de NRMSE, R2 y captura de varianza held-out para generadores compartidos de una a seis dimensiones.",
    )
    add_body(
        doc,
        f"La acción espacial conocida sobre mu también tuvo residuos altos: medianas {rotation_decisions['spatial_action']['normal_vae']['median']:.3f} normal y {rotation_decisions['spatial_action']['so2_vae']['median']:.3f} SO(2), con 0/25 por debajo del umbral 0,25. Su canonicalización redujo la varianza intraórbita a cerca de 0,34 de la identidad, pero aumentó la varianza entre parches unas 4,7--4,8 veces; W/B quedó en 3,235 normal y 3,344 SO(2). Como esa acción espacial sobre mu no fue verificada, la reducción no respalda separación contenido--pose. Con la acción aprendida d=1, canonicalizar equivale exactamente a identidad y los probes de ángulo no mejoran. H6 queda no resuelta —no apoyada—, no rechazada.",
    )
    add_figure(
        doc,
        rotation_figures / "07-canonicalization-factorization.png",
        6.45,
        "Figura 8g. Canonicalización mediante la acción espacial conocida y la acción reducida aprendida. La conocida no satisface el residuo de equivarianza y la aprendida coincide con identidad; no hay evidencia válida de factorización local.",
        "Comparación de varianza intraórbita, varianza entre parches, recuperación de identidad y fuga angular tras canonicalización.",
    )

    add_heading(doc, "Interpretación geométrica", 2)
    add_table(
        doc,
        ["Nivel", "Qué significa aquí", "Resultado"],
        [
            [
                "Órbita por parche",
                "Curva de mu al rotar una imagen",
                "Continua tras reparar el operador; evidencia descriptiva",
            ],
            [
                "Acción espacial C4 prescrita",
                "La arquitectura prescribe rho; se prueba su conmutación con el decoder",
                f"Verificada empíricamente en {FIXED25_COUNT}/{FIXED25_COUNT} para R90, R180 y R270",
            ],
            [
                "Acción reducida aprendida",
                "Un mismo exp(theta A) predice parches y ángulos nuevos",
                "No demostrada",
            ],
            [
                "Factorización local",
                "Q(-theta) estabiliza contenido sin borrar diferencias",
                "No apoyada; no resuelta sin acción reducida válida",
            ],
            [
                "Producto global",
                "Un SO(2) x R^n único sobre toda la representación",
                "Fuera de alcance e imposible de concluir con 25 parches",
            ],
        ],
        widths=[1.4, 3.55, 1.75],
        font_size=8.2,
    )
    add_body(
        doc,
        f"Prescrito por arquitectura: existen tipos F0/F1 y una acción espacial candidata rho. Verificado empíricamente: el operador de entrada corregido y la conmutación casi exacta del decodificador del checkpoint SO(2) con C4 en {FIXED25_COUNT}/{FIXED25_COUNT}; no la equivarianza continua del encoder. Observado en el checkpoint entrenado, sin atribución causal: la ventaja comparativa en consistencia de rutas decodificadas para ángulos interpolados, junto con órbitas crudas complejas y de alta dimensionalidad. Transferencia entre parches: la acción espacial C4 no usa parámetros por parche, pero ninguna acción reducida aprendida superó identidad. Lectura tipo fibrado: la canonicalización decodificada bajo giros C4 exactos fue fuerte, pero no exacta, y no define un cociente de contenido ni fases locales fiables; una interpretación de fibras, gauges o topología permanece especulativa.",
    )

    doc.add_page_break()
    add_heading(doc, "9. Visualizacion PCA del espacio latente", 1)
    add_body(
        doc,
        "La visualizacion tipo EQ-VAE proyecta descriptores latentes espaciales a tres componentes principales y los presenta como falso color. Su proposito es mostrar organizacion espacial dentro del latente; PCA es un sistema de coordenadas, no una metrica de calidad ni de equivarianza.",
    )
    add_figure(
        doc,
        legacy_orbit_figures / "05-paper-style-latent-pca.png",
        6.45,
        "Figura 9. Visualizacion PCA-RGB de cuatro parches mostrados (ranks 0, 8, 12 y 24) de la fuente fixed-25. Los modelos usan ajustes PCA separados; los colores y rangos no deben compararse como magnitudes absolutas entre modelos.",
        "Visualizacion PCA en falso color compara mapas latentes espaciales del modelo normal y del modelo SO2.",
    )
    add_body(
        doc,
        "La coherencia local mediana r_edge fue 1,4567 para el modelo normal y 1,4676 para SO(2); menor es mas coherente bajo este proxy. La diferencia no respalda una ventaja de coherencia local para SO(2). Un lector RGB local exploratorio tambien obtuvo R2=0,286 para normal y R2=0,189 para SO(2). Ninguno de estos diagnosticos es un endpoint confirmatorio.",
    )

    doc.add_page_break()
    add_heading(doc, "10. Diagnostico WSI en el test sellado", 1)
    nm, sm = mil_normal["metrics"], mil_so2["metrics"]
    add_heading(doc, "Desarrollo supervisado y seleccion", 2)
    add_body(
        doc,
        "Los clasificadores se ajustaron después de congelar cada representación. En MIL se usaron 106 WSI de entrenamiento y 23 WSI de validación; ambos entrenamientos terminaron por parada temprana. La pérdida de train se registró en línea antes de cada actualización y se resume por ventanas de 53 actualizaciones, una WSI por actualización, mientras que el F1 y la entropía cruzada de validación se calcularon con el modelo fijo en cada frontera. Para tejido se entrenó un clasificador por representación y presupuesto con subconjuntos balanceados. La validación sirvió para seleccionar checkpoints y no se reutiliza como evidencia del test sellado.",
    )
    add_figure(
        doc,
        supervised_development_chart,
        6.9,
        "Figura 10. Desarrollo supervisado. Arriba izquierda: F1 macro de validación MIL; los puntos resaltan los checkpoints seleccionados (0,448 normal en la actualización 3.127 y 0,591 SO(2) en 5.671). Arriba derecha: entropía cruzada de train registrada en línea y promediada por cada ventana de 53 actualizaciones, junto con la entropía cruzada de validación en cada frontera. Como el modelo cambia dentro de cada ventana, la curva de train diagnostica la optimización y no equivale a evaluar un checkpoint fijo sobre todo train. Abajo: mejor F1 macro de validación tisular por presupuesto. Ningún panel es test sellado ni variación entre semillas.",
        "Tres paneles muestran F1 macro de validacion WSI, entropia cruzada de train y validacion WSI, y el mejor F1 macro de validacion tisular por presupuesto.",
    )
    add_heading(doc, "Resultado WSI sellado", 2)
    add_body(
        doc,
        "Se entreno una unica trayectoria de clasificador MIL por representacion y se evaluo una sola vez sobre 23 WSI sellados. La tarea contiene cinco diagnosticos; LGSC y MC tienen solo dos WSI cada uno, por lo que las conclusiones por clase son fragiles.",
    )
    add_table(
        doc,
        ["Metrica", "Normal", "SO(2)", "Normal - SO(2) [IC 95%]"],
        [
            [
                "F1 macro",
                f"{nm['macro_f1']:.3f}",
                f"{sm['macro_f1']:.3f}",
                f"{mil_pair['macro_f1']['difference']:+.3f} [{mil_pair['macro_f1']['confidence_low']:+.3f}; {mil_pair['macro_f1']['confidence_high']:+.3f}]",
            ],
            [
                "Exactitud balanceada",
                f"{nm['balanced_accuracy']:.3f}",
                f"{sm['balanced_accuracy']:.3f}",
                f"{mil_pair['balanced_accuracy']['difference']:+.3f} [{mil_pair['balanced_accuracy']['confidence_low']:+.3f}; {mil_pair['balanced_accuracy']['confidence_high']:+.3f}]",
            ],
            [
                "Exactitud",
                f"{nm['accuracy']:.3f}",
                f"{sm['accuracy']:.3f}",
                f"{mil_pair['accuracy']['difference']:+.3f} [{mil_pair['accuracy']['confidence_low']:+.3f}; {mil_pair['accuracy']['confidence_high']:+.3f}]",
            ],
            [
                "Entropia cruzada",
                f"{nm['mean_ce']:.3f}",
                f"{sm['mean_ce']:.3f}",
                f"{mil_pair['mean_ce']['difference']:+.3f} [{mil_pair['mean_ce']['confidence_low']:+.3f}; {mil_pair['mean_ce']['confidence_high']:+.3f}]",
            ],
        ],
        widths=[1.65, 1.0, 1.0, 3.05],
        font_size=8.5,
    )
    add_figure(
        doc,
        mil_chart,
        6.55,
        "Figura 11. Metricas de diagnostico WSI en el test sellado (n=23 WSI). SO(2) es numericamente mayor en F1 macro y exactitud; la diferencia de F1 macro no está establecida porque su IC 95% incluye cero.",
        "Grafica de barras de F1 macro, exactitud balanceada y exactitud para 23 WSI bajo ambos modelos.",
    )
    add_body(
        doc,
        "La entropia cruzada, endpoint secundario, fue menor para SO(2) y su intervalo pareado favorece a SO(2). Este resultado debe interpretarse junto con el unico entrenamiento del clasificador, el test pequeno y el bajo soporte de dos clases.",
    )

    doc.add_page_break()
    add_heading(doc, "Desglose por clase en diagnostico WSI", 2)
    add_body(
        doc,
        "Las matrices permiten ver que la diferencia global no se distribuye uniformemente. SO(2) obtuvo F1 mayor en CC, EC, HGSC y LGSC, mientras que el modelo normal obtuvo F1 mayor en MC. En particular, el VAE normal no recupero casos EC ni LGSC y SO(2) no recupero MC. Estas observaciones son descriptivas: los soportes por clase son CC=5, EC=6, HGSC=8, LGSC=2 y MC=2 WSI.",
    )
    add_figure(
        doc,
        mil_class_chart,
        6.45,
        "Figura 12. Matrices de confusion normalizadas por clase real y F1 por diagnostico en el test sellado. Cada celda incluye conteo y porcentaje dentro de su fila; los soportes 5/6/8/2/2 muestran por que LGSC y MC no permiten conclusiones estables por clase.",
        "Dos matrices de confusion y barras de F1 por clase comparan los clasificadores WSI normal y SO2 en 23 laminas.",
    )
    add_heading(doc, "Atribucion espacial a parches", 2)
    add_body(
        doc,
        "No se incluyo un mapa de atencion por WSI porque la corrida sellada almaceno logits, prediccion, tamano de bolsa e identidad y resumen del grafo, pero no compuertas por parche. La atencion local describe intercambio entre vecinos; las compuertas sigmoidales del resumen global no son probabilidades ni atribuciones causales o especificas de clase. Generar un mapa exige una nueva inferencia instrumentada sobre latentes y pesos congelados y, para superponerlo sobre tejido real, acceso al thumbnail o WSI. Un analisis futuro deberia preseleccionar el WSI antes de inspeccionar el mapa, usar la misma escala espacial para ambos modelos y etiquetar el resultado como intensidad de lectura interna o sensibilidad post-hoc, no como explicacion causal.",
    )

    doc.add_page_break()
    add_heading(doc, "11. Eficiencia de etiquetas en tejido", 1)
    add_body(
        doc,
        "La prueba de tejido uso 31.572 parches de 23 WSI: 21.796 tumorales, 8.500 de estroma y 1.276 de necrosis. La necrosis aparece en solo cinco WSI. Se compararon cinco presupuestos de etiquetas por clase con intervalos simultaneos para controlar la familia de contrastes. El bootstrap se estratifico por soporte tisular observado: 2 WSI con solo tumor, 16 con tumor y estroma, y 5 con los tres tejidos.",
    )
    tissue_rows = []
    for row in tissue["rows"]:
        sim = row["primary_simultaneous_95"]
        tissue_rows.append([
            f"{row['labels_per_class']:,}".replace(",", "."),
            f"{row['normal_vae']['metrics']['macro_f1']:.3f}",
            f"{row['so2_vae']['metrics']['macro_f1']:.3f}",
            f"{sim['estimate']:+.3f} [{sim['lower']:+.3f}; {sim['upper']:+.3f}]",
        ])
    add_table(
        doc,
        [
            "Etiquetas/clase",
            "F1 macro normal",
            "F1 macro SO(2)",
            "Normal - SO(2), IC simultaneo 95%",
        ],
        tissue_rows,
        widths=[1.3, 1.35, 1.35, 2.7],
        font_size=8.4,
    )
    add_figure(
        doc,
        tissue_chart,
        6.55,
        "Figura 13. Curvas de F1 macro por presupuesto en el test sellado de tejido: 31.572 parches de 23 WSI. Solo 500 etiquetas por clase presenta un intervalo simultaneo que excluye cero y favorece a SO(2).",
        "Grafica lineal de F1 macro a cinco presupuestos de etiquetas por clase para ambos modelos; 500 está marcado como diferencia establecida.",
    )
    aulc = tissue["aulc"]
    add_body(
        doc,
        f"Resumen global. La AULC trapezoidal normalizada sobre log10(etiquetas por clase), sin suavizado, fue {aulc['normal_vae']['estimate']:.3f} para normal y {aulc['so2_vae']['estimate']:.3f} para SO(2). Diferencia normal menos SO(2) = {aulc['primary_simultaneous_95']['estimate']:+.3f}; IC simultaneo 95% [{aulc['primary_simultaneous_95']['lower']:+.3f}; {aulc['primary_simultaneous_95']['upper']:+.3f}]. El intervalo incluye cero: la ventaja global es inconclusa. Esta AULC resume el F1 macro medio en el rango logaritmico preespecificado; no demuestra que un modelo necesite menos etiquetas.",
        bold_prefix="Resumen global.",
    )
    add_figure(
        doc,
        tissue_class_chart,
        6.55,
        "Figura 14. F1 por tipo de tejido a cada presupuesto de etiquetas. Son estimaciones puntuales descriptivas; la inferencia primaria simultanea se definio para F1 macro, no para seleccionar retrospectivamente una clase. Necrosis tiene 1.276 parches pero aparece en solo cinco WSI.",
        "Tres curvas comparan F1 de tumor, estroma y necrosis a cinco presupuestos para los modelos normal y SO2.",
    )
    add_body(
        doc,
        "El desglose muestra que las diferencias cambian segun tejido y presupuesto. SO(2) presenta su senal mas consistente en tumor a presupuestos bajos; estroma y, especialmente, necrosis fluctuan mas. En el presupuesto maximo, la ventaja de F1 macro del modelo normal se explica en parte por un F1 de necrosis mayor. Esta lectura no convierte las curvas por clase en contrastes confirmatorios.",
    )

    doc.add_page_break()
    add_heading(doc, "12. Sintesis integrada", 1)
    add_table(
        doc,
        ["Dimension", "Senal observada", "Conclusion permitida"],
        [
            [
                "Fidelidad de reconstruccion",
                "Normal ligeramente mejor",
                "Sin diferencia primaria establecida",
            ],
            [
                "Diagnostico WSI",
                "SO(2) mayor en F1 macro y exactitud",
                "Senal favorable, no superioridad general",
            ],
            [
                "Tejido con pocas etiquetas",
                "SO(2) mayor en 250-2.500",
                "Evidencia puntual a 500; AULC inconclusa",
            ],
            [
                "Presupuesto maximo de tejido",
                "Normal mayor en F1 macro",
                "El beneficio de SO(2) no persiste uniformemente",
            ],
            [
                "Geometría cruda de mu",
                "Operador corregido; H1--H4 de Spec 0050 no pasan",
                "Sin acción reducida aprendida ni factorización demostrada",
            ],
            [
                "Rutas decodificadas",
                f"SO(2) menor en {decoded_action_review['so2_favorable_patches']}/{FIXED25_COUNT}; decoder C4 casi sin error",
                "Ventaja comparativa de rutas y acción C4 del decoder; no SO(2) continuo",
            ],
            [
                "Eficiencia parametrica",
                "SO(2) usa 29,8% de parametros",
                "Ventaja de compactacion; confusor de capacidad",
            ],
        ],
        widths=[1.6, 2.15, 3.0],
        font_size=8.4,
    )
    add_heading(doc, "Mensaje central para el profesor", 2)
    add_body(
        doc,
        "La arquitectura continua SO(2) no mejoró automáticamente la reconstrucción ni mostró una acción reducida aprendida o una factorización demostrable en el posterior final. En el checkpoint SO(2) se observó una acción espacial C4 casi exacta al transformar el embedding antes del decoder y una reducción comparativa consistente del error entre rutas decodificadas para ángulos interpolados. El checkpoint también es mucho más compacto y presenta señales de utilidad en diagnóstico WSI y ciertos regímenes de bajo etiquetado; este estudio no aísla causalmente la simetría de la capacidad u otros factores. La conclusión responsable es específica por endpoint: se verificó empíricamente un mecanismo de rotación C4 en el decoder, mientras que la geometría cruda de mu y la separación contenido--pose siguen sin validarse.",
    )

    add_heading(doc, "13. Limitaciones", 1)
    add_bullet(
        doc,
        "Una sola trayectoria de entrenamiento por VAE y una sola trayectoria de clasificador por presupuesto: los intervalos no incluyen variacion entre semillas de entrenamiento.",
    )
    add_bullet(
        doc,
        "El test WSI tiene 23 laminas; LGSC y MC solo aportan dos WSI cada una.",
    )
    add_bullet(
        doc,
        "El test de tejido contiene 31.572 parches, pero su incertidumbre efectiva depende de 23 WSI y la necrosis aparece en solo cinco.",
    )
    add_bullet(
        doc,
        "Los intervalos de tejido condicionan la inferencia a los estratos de soporte observados (2/16/5 WSI). Los cinco presupuestos comparten parches de test y datos de entrenamiento anidados, por lo que sus resultados estan correlacionados.",
    )
    add_bullet(
        doc,
        "Los intervalos inferenciales sellados remuestrean WSI y no cubren selección de parches, arquitectura ni nuevas semillas. En la geometría post-hoc, Spec 0050 remuestrea los 25 parches pareados; Spec 0051 resume primero las diferencias por sus 16 WSI de origen y remuestrea esas 16 WSI. Ambos intervalos son descriptivos y no poblacionales.",
    )
    add_bullet(
        doc,
        "Las desviaciones estandar por parche son dispersion descriptiva, no error estandar ni intervalo de confianza.",
    )
    add_bullet(
        doc,
        "SO(2) tiene aproximadamente 70,2% menos parametros; simetria y capacidad no estan aisladas como causas separadas.",
    )
    add_bullet(
        doc,
        "La geometría rotacional es post-hoc y exploratoria sobre 25 parches fijos de validación. No incorpora test sellado, variación entre semillas ni incertidumbre de una población nueva.",
    )
    add_bullet(
        doc,
        "La ventaja decodificada de Spec 0051 compara rutas dentro de cada modelo. Sus objetivos son transformaciones de la reconstrucción del propio modelo, no imágenes originales ni una nueva evaluación de fidelidad.",
    )
    add_bullet(
        doc,
        "La acción casi exacta se verificó en rotaciones de cuarto de vuelta alineadas con la cuadrícula. Los ángulos intermedios requieren interpolación y no sostienen la misma afirmación de exactitud; tampoco se evaluaron rotaciones continuas fuera de la rejilla de 5 grados.",
    )
    add_bullet(
        doc,
        "La interpolación bilineal impone parte de la regularidad de la trayectoria de entrada; las conclusiones dependen de controlar esa ruta y de la resolución angular. El efecto a 1 grado se invirtió a 5 grados.",
    )
    doc.add_page_break()
    add_heading(doc, "13. Limitaciones (continuación)", 2)
    add_bullet(
        doc,
        "PC1--PC6 capturaron alrededor de 11% de la variación orbital. Los generadores de d=1,...,6 examinan sólo una fracción pequeña del posterior espacial y no excluyen acciones compartidas de mayor dimensión o no lineales.",
    )
    add_bullet(
        doc,
        "La dimensión d=1 se eligió mediante una regla heurística de un error estándar leave-one-patch-out; los pliegues están correlacionados y la tolerancia no constituye incertidumbre poblacional.",
    )
    add_bullet(
        doc,
        "El addendum de revisión conserva ocho resultados por parche para el probe angular de factorización; aun así, la acción aprendida es identidad y los prerrequisitos no degenerados fallan, por lo que el estado es no resuelto.",
    )
    add_bullet(
        doc,
        "La evaluacion WSI sellada no guardo compuertas por parche. Cualquier mapa espacial de lectura o atribucion requiere una nueva inferencia instrumentada y una interpretacion post-hoc separada.",
    )

    doc.add_page_break()
    add_heading(doc, "14. Conclusiones para la presentacion", 1)
    add_bullet(
        doc,
        "Los requerimientos de evaluacion del profesor quedaron cubiertos para ambos modelos.",
    )
    add_bullet(
        doc,
        "La reconstruccion sellada no demuestra una diferencia primaria; el modelo normal conserva una ligera ventaja descriptiva.",
    )
    add_bullet(
        doc,
        "SO(2) ofrece compactacion parametrica fuerte y senales favorables en tareas posteriores, especialmente con 500 etiquetas por clase y en diagnostico WSI, pero con incertidumbre y soporte limitado.",
    )
    add_bullet(
        doc,
        "El barrido corregido retira la anterior conclusión 25/25 como evidencia general: la diferencia descriptiva a 1 grado no alcanzó el umbral fijado y se invirtió a 5 grados.",
    )
    add_bullet(
        doc,
        f"En las rutas decodificadas, SO(2) mostró una ventaja comparativa consistente para el barrido de 5 grados: menor error de acción y canonicalización en {decoded_action_review['so2_favorable_patches']}/{FIXED25_COUNT} parches. Sus medianas {es_decimal(float(decoded_action_review['so2']), 3)} y {es_decimal(float(decoded_canonical_review['so2']), 3)} no cumplieron el criterio funcional absoluto <=0,50, que es distinto de la comparación entre modelos.",
    )
    add_bullet(
        doc,
        f"Para R90, R180 y R270 exactas se verificó empíricamente una acción espacial C4 compartida del decoder del checkpoint SO(2), con razón de acción agregada {es_decimal(float(decoded_action_quarter['so2']), 8)}. El resultado no se extiende automáticamente a SO(2) continuo, al encoder ni al grupo D4 completo.",
    )
    add_bullet(
        doc,
        "No se demostró una acción rotacional reducida aprendida; la factorización local contenido--pose queda no resuelta y no se respalda una interpretación de producto global, fibrado o gauge con esta muestra.",
    )
    add_bullet(
        doc,
        "La contribución del experimento es mostrar dónde el checkpoint SO(2) presentó señales favorables, dónde no y dónde aparece una posible compensación descriptiva con la fidelidad de reconstrucción, sin atribución causal a un único componente arquitectónico.",
    )
    doc.add_page_break()
    add_heading(doc, "Anexo A. Cobertura de lo solicitado por el profesor", 1)
    coverage_table = add_table(
        doc,
        ["Solicitud", "Donde se presenta", "Estado"],
        [
            ["Graficas de resultados de la linea base", "Figuras 1-5", "Completo"],
            ["MAE, MSE, PSNR, SSIM", "Secciones 4 y 6", "Completo"],
            ["Desviacion estandar, n y boxplots", "Tablas y Figuras 2 y 4", "Completo"],
            ["25 originales y reconstrucciones", "Figura 3", "Completo"],
            ["Rotaciones 90/180/270", "Figura 6", "Completo"],
            [
                "Órbita latente 0--359 y geometría local",
                "Figura 7",
                "Corregido; evidencia previa supersedida",
            ],
            [
                "Órbitas de los 25 parches cada 1 grado",
                "Figuras 8a--8c",
                "Corregido con controles de entrada",
            ],
            [
                "Armónicos y 48 campos F1",
                "Figuras 8d y 8e",
                "Completo; hipótesis no respaldadas",
            ],
            [
                "Acción compartida y factorización",
                "Figuras 8f y 8g",
                "Completo; acción reducida no demostrada",
            ],
            [
                "Rotar y desrotar el embedding decodificado",
                "Figuras 6a y 6b; Sección 8",
                "Completo; ventaja pareada y C4 exacto",
            ],
            [
                "Rotaciones y flips exactos D4",
                "Figura 6b; Sección 8",
                "Completo; sólo flip diagonal pasa en SO(2)",
            ],
            ["Visualizacion PCA espacial del latente", "Figura 9", "Completo"],
            [
                "Repetir evaluacion para SO(2)",
                "Todas las tablas y figuras comparativas",
                "Completo",
            ],
            [
                "Validar componentes SO(2)",
                "Seccion 2: no linealidades, normalizacion, upsampling, mu/logvar/z y runtime",
                "Completo",
            ],
            [
                "Barrido continuo 0-359 grados",
                "Sección 8: artefacto fijo-25 corregido",
                "Completo; validación post-hoc",
            ],
            [
                "Utilidad posterior y desglose por clase",
                "Figuras 10-14; Secciones 10 y 11",
                "Completo",
            ],
            [
                "Atribucion espacial WSI",
                "Seccion 10: frontera y protocolo futuro",
                "Pendiente; requiere nueva inferencia",
            ],
        ],
        widths=[2.55, 3.15, 1.0],
        font_size=8.3,
    )

    doc.add_page_break()
    add_heading(doc, "Anexo B. Definicion breve de metricas", 1)
    metrics_table = add_table(
        doc,
        ["Metrica", "Lectura"],
        [
            [
                "MAE",
                "Error absoluto medio en dominio normalizado [-1, 1]; menor es mejor.",
            ],
            [
                "MSE",
                "Error cuadratico medio en dominio normalizado [-1, 1]; menor es mejor.",
            ],
            [
                "PSNR",
                "Relacion senal-ruido pico en dB, imagen recortada a [0, 1]; mayor es mejor.",
            ],
            ["SSIM", "Similitud estructural en imagen [0, 1]; mayor es mejor."],
            [
                "F1 macro",
                "Promedio no ponderado del F1 de cada clase; valora por igual clases frecuentes y raras.",
            ],
            [
                "Exactitud balanceada",
                "Promedio del recall por clase; reduce el predominio de las clases frecuentes.",
            ],
            [
                "Entropia cruzada",
                "Perdida probabilistica; menor es mejor y penaliza predicciones confiadas incorrectas.",
            ],
            [
                "AULC",
                "Area trapezoidal normalizada del F1 macro sobre log10(etiquetas por clase), sin suavizado; equivale al rendimiento medio en el rango logaritmico preespecificado y no demuestra necesitar menos etiquetas.",
            ],
            [
                "Razon de linealidad local",
                "RMS cíclico de la segunda diferencia de mu dividido por el RMS cíclico de la primera; depende del paso angular y debe compararse con la trayectoria de entrada.",
            ],
            [
                "CV del paso angular",
                "Desviacion estandar dividida por la media de los pasos RMS ciclicos de mu entre angulos consecutivos; menor indica velocidad de recorrido mas uniforme.",
            ],
            [
                "Concentración armónica",
                "Fracción de la energía orbital cruda total capturada por m=1,...,6 dentro de PC1--PC6; una frecuencia limpia requiere además transferencia y suficiente varianza capturada.",
            ],
            [
                "NRMSE / R2 del generador",
                "Predicción de ángulos y parches retenidos con una única acción exp(theta A); NRMSE menor y R2 mayor son mejores, siempre contra identidad y ángulos barajados.",
            ],
            [
                "Residuo de acción conocida",
                "Norma relativa entre E(R_theta x) y la rotación espacial de E(x) dentro del disco latente; cero representa equivarianza exacta.",
            ],
            [
                "Razón de acción decodificada",
                "RMS entre D(rho_T mu(x)) y T D(mu(x)), dividido por el error de dejar mu sin transformar. Uno es la línea base de no actuar; menor es mejor y 0,50 significa reducir ese error al menos a la mitad.",
            ],
            [
                "Razón de canonicalización decodificada",
                "RMS entre D(rho_T^-1 mu(Tx)) y D(mu(x)), dividido por el error sin desrotar el embedding. Mide consistencia entre rutas, no fidelidad frente a la imagen original.",
            ],
            [
                "F = W / B",
                "Varianza intraórbita después de canonicalizar dividida por varianza entre parches; sólo apoya factorización si la acción es válida, W disminuye y el contenido se conserva.",
            ],
        ],
        widths=[1.5, 5.2],
        font_size=8.3,
    )
    for table in (coverage_table, metrics_table):
        for row in table.rows:
            for cell in row.cells:
                set_cell_margins(cell, top=45, bottom=45)

    add_heading(doc, "Trazabilidad de resultados", 2)
    add_body(
        doc,
        "Fuentes aceptadas: professor_metrics_v1; vae_test_reconstruction_scored_v1; ubc_ocean_mil_test_scored_v1; tissue_test_scored_v1; la PCA espacial fija de frozen_vae_rotation_orbits; Kaggle maximshtefan/eqvae-corrected-rotation-geometry/1; y Kaggle maximshtefan/eqvae-decoded-latent-transform-audit/1. Para Spec 0050, el resumen geométrico tiene SHA-256 aab68ca5b7cbe989ea1cd79740a64d83784441f85ffde09a6ee6ebebe7faf43c y el addendum ec74b37f509205bca12294321844a3cfb48dbd2f1763c8ff94290fe064e86f94. Para Spec 0051, el resumen tiene SHA-256 616e3a6c9caa6b8e72a0b21f41a5121d7a5baa420f8023b35775fedaf252a754, el manifiesto a1168b419080305138f7aedca4c1ed345c3d5947e78299b0e5ea1430a75d3c50 y el addendum b0fa0cc2eebe1207fcc2da7a00c30b9b85c40f4353e332131423f76f6c380e34. El artefacto dense-rotation-population/1 se conserva sólo como procedencia supersedida. El informe no recalcula inferencia, cambia umbrales ni selecciona ejemplos.",
    )

    doc.save(output_docx)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    default_repo = Path(__file__).resolve().parents[1]
    parser.add_argument("--repo-root", type=Path, default=default_repo)
    parser.add_argument("--output-docx", type=Path, default=None)
    parser.add_argument("--work-dir", type=Path, default=None)
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    repo = args.repo_root.resolve()
    output = (
        args.output_docx
        or repo / "reports/professor/informe_final_experimento_eqvae.docx"
    )
    work_dir = args.work_dir or repo / ".agent_tmp/professor_report_build"
    build_report(repo, output.resolve(), work_dir.resolve())
    print(output.resolve())


if __name__ == "__main__":
    main()
